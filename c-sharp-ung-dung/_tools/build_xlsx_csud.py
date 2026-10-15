import os, re, sys, unicodedata
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

OUT = sys.argv[1]
REPO = "/Users/tuananhcan/Downloads/programming-course-demos"
COURSE_DIR = "c-sharp-ung-dung"
BASE_URL = "https://codegym-vn.github.io/programming-course-demos/"


def slug(s):
    s = s.replace("C#", "c-sharp").replace("đ", "d").replace("Đ", "D")
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn").lower()
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


MODULES = [
    "Cấu trúc dữ liệu nâng cao: Collections",
    "Xử lý lỗi và Ngoại lệ (Exception Handling) Cơ bản",
    "Xử lý lỗi và Ngoại lệ (Exception Handling) Nâng cao",
    "Đọc và Ghi File Văn bản",
    "Làm việc với Dữ liệu có Cấu trúc (CSV)",
    "Lưu trữ Đối tượng với Serialization",
]

N, T, K = "Nên làm", "Làm thêm nếu muốn", "Không cần"
R, A = "Bài đọc", "Hoạt động (đã có nội dung)"

# Thư viện thành phần hình minh họa dùng chung
COMPS = [
    ("H01", "Dãy ô List<T>", "Hàng ô có chỉ số bên dưới; ô đã dùng nền đặc, ô dự phòng (Capacity) viền nét đứt; nhãn Count và Capacity. Khi Insert/Remove, các ô phía sau trượt sang phải/trái (mũi tên dịch chỉ số), ô mới nền vàng, ô bị xóa mờ dần."),
    ("H02", "Tủ ngăn khóa – giá trị", "Mỗi phần tử Dictionary là một ngăn kéo có nhãn khóa ở mặt trước và giá trị bên trong. Tra cứu: mũi tên đi thẳng tới ngăn; ngăn mới trượt vào (xanh), ngăn cập nhật nháy vàng, ngăn bị xóa mờ đỏ."),
    ("H03", "Thẻ Exception", "Thẻ viền đỏ: biểu tượng cảnh báo, tên loại (FormatException...), thông điệp thật của .NET, dòng code gây lỗi, một dòng 'Cách phòng tránh'. Dùng thống nhất cho mọi bài về lỗi."),
    ("H04", "Sơ đồ luồng try / catch / finally", "Các khối chữ nhật try, catch, finally, 'sau khối' nối bằng mũi tên. Đường đi thực tế được tô (xanh: bình thường, đỏ: nhánh lỗi). Dòng code bị bỏ qua trong try hiện mờ kèm nhãn 'không chạy'."),
    ("H05", "Bộ lọc catch xếp dọc", "Các catch xếp chồng như những cửa lọc. Exception (viên đỏ) rơi từ trên xuống: cửa không khớp hiện ✗ xám, cửa khớp hiện ✓ xanh và mở ra; các cửa phía dưới không được xét."),
    ("H06", "Cây phân cấp lớp", "Cây thụt lề hoặc sơ đồ nút: Exception → SystemException → ... Khi so khớp, tô đường đi từ lớp con lên lớp cha để thấy quan hệ 'là một'. Nút mới (exception tự tạo) được vẽ thêm dần."),
    ("H07", "Ngăn xếp lời gọi (call stack)", "Các khung hàm chồng lên nhau (Main ở đáy). Exception là quả bóng đỏ đi lên từng khung; khung không có catch phù hợp bị đóng lại (mờ), khung có catch bắt lấy (viền xanh)."),
    ("H08", "Tài nguyên có trạng thái + bộ đếm", "Biểu tượng tệp/kết nối với đèn trạng thái (Đóng xám, Đang mở xanh, Rò rỉ đỏ) và bộ đếm 'Đang mở: n'. Dùng để thấy tài nguyên có được giải phóng hay không."),
    ("H09", "Đèn kiểm tra / băng chuyền", "Dải trạm kiểm tra nằm ngang; mỗi trạm là một đèn: xanh khi đạt, đỏ khi trượt, các trạm sau tắt. Biến thể băng chuyền: 'kiện hàng' dữ liệu đi qua từng trạm xử lý."),
    ("H10", "RAM và Ổ đĩa", "Hai khung hình: 'Bộ nhớ chương trình (RAM)' và 'Ổ đĩa' có biểu tượng tệp. Nút trạng thái chương trình: Đang chạy → Đã tắt (RAM xóa trắng) → Chạy lại."),
    ("H11", "Ống dẫn Stream", "Ống nằm ngang nối Tệp và Chương trình; từng byte (ô nhỏ ghi ký tự + mã số) chảy qua ống; con trỏ Position chạy trên dải byte của tệp; nhãn Length."),
    ("H12", "Trình xem tệp", "Khung giống Notepad: số dòng, nội dung, ký hiệu ↵ cho xuống dòng; con trỏ đọc ▶ hoặc con trỏ ghi |; chế độ 'trước / sau' đặt cạnh nhau, dòng mới tô xanh, dòng bị ghi đè gạch ngang."),
    ("H13", "Bộ đệm (buffer)", "Ba khối Chương trình → Bộ đệm → Tệp trên đĩa. Bộ đệm là hộp có vạch đầy dần; khi Flush/Close, nội dung 'đổ' xuống tệp và kích thước tệp tăng."),
    ("H14", "Chip đường dẫn", "Đường dẫn tách thành các chip: ổ đĩa | thư mục | thư mục | tên tệp | phần mở rộng. Mỗi phương thức của Path tô sáng các chip mà nó trả về."),
    ("H15", "Cây thư mục", "Cây có biểu tượng thư mục/tệp và thụt lề; nút mới tạo tô xanh, nút bị xóa mờ đỏ, nút được liệt kê viền vàng, vị trí gây lỗi gắn biểu tượng ổ khóa hoặc dấu chấm than."),
    ("H16", "Bảng tính ↔ văn bản CSV", "Trái: lưới ô kiểu Excel (hàng tiêu đề đậm). Phải: văn bản CSV thô. Mỗi bước tô cùng màu một hàng ở hai bên; dấu phẩy được tô như đường kẻ ô."),
    ("H17", "Dây chuyền dòng ↔ đối tượng", "Dây chuyền ngang: dòng văn bản → Split/Join → mảng ô → chuyển kiểu → thẻ đối tượng → List. Mỗi bản ghi lần lượt đi qua từng trạm."),
    ("H18", "Bộ phân tích theo trạng thái", "Con trỏ chạy trên từng ký tự của một dòng; huy hiệu trạng thái 'Trong nháy: Có/Không'; ký tự bị bỏ qua tô xám; trường hoàn tất được thả vào ô kết quả."),
    ("H19", "Thẻ đối tượng ↔ JSON/XML", "Thẻ đối tượng bên trái, văn bản JSON/XML bên phải; đường nối màu từ mỗi thuộc tính tới phần tương ứng. JSON tô màu theo kiểu: chuỗi xanh lá, số cam, bool tím, null xám."),
    ("H20", "Dải byte nhị phân", "Các ô hex (16 ô/hàng) đầy dần khi ghi; mỗi trường một màu; chú thích 'Id: 4 byte', 'Name: 1 byte độ dài + 4 byte'..."),
    ("H21", "Hai thùng phân loại", "Hai thùng 'Hợp lệ' (thẻ xanh) và 'Lỗi' (thẻ đỏ ghi số dòng + lý do); mỗi bản ghi sau khi kiểm tra được thả vào đúng thùng."),
    ("H22", "Bảng điểm so sánh 2 cột", "Thanh đếm 'số bước / số lần so sánh / số chỗ phải sửa' dưới mỗi cột, cập nhật mỗi bước; bước cuối gắn cờ 'phù hợp hơn' cho cột thắng."),
]

ROWS = []


def row(m, lesson, act, typ, lvl, goal="", form="", visual="", comps="", tabs="", steps=None, data="", key="", cs="", note=""):
    ROWS.append(dict(m=m, lesson=lesson, act=act, typ=typ, lvl=lvl, goal=goal, form=form, visual=visual,
                     comps=comps, tabs=tabs, steps="\n".join(steps or []), data=data, key=key, cs=cs, note=note))


# ======================= CHƯƠNG 1: COLLECTIONS =======================
L = "Giới thiệu về Collections và List<T> cơ bản"
row(1, L, "Giới thiệu Collections trong C#", R, N,
    goal="Thấy vì sao cần List<T>: mảng có kích thước cố định, còn List<T> tự mở rộng khi thêm phần tử.",
    form="2 cột song song: Mảng int[] | List<int>",
    visual="Mỗi cột một dãy ô có chỉ số. Ô dự phòng của List vẽ nét đứt. Khi mảng bị ghi quá giới hạn, ô 'không tồn tại' nháy đỏ. Khi List hết chỗ, hiện hiệu ứng cấp dãy ô mới gấp đôi và chép phần tử sang.",
    comps="H01, H03, H22", tabs="1 tab: Thêm phần tử thứ 4 và thứ 5",
    steps=["Hai cột cùng cần lưu điểm 8, 6, 9. Trái: int[] a = new int[3]. Phải: List<int> l = new List<int>() (Count 0, Capacity 0).",
           "Trái gán a[0..2]. Phải gọi Add 3 lần: Count 1→3, lần Add đầu cấp sẵn 4 chỗ (Capacity 4, ô thứ 4 nét đứt).",
           "Cả hai muốn thêm điểm 7.",
           "Trái: a[3] = 7 → không có ô chỉ số 3, console hiện IndexOutOfRangeException, chương trình dừng.",
           "Phải: l.Add(7) → lấp ô dự phòng, Count 4, Capacity 4.",
           "Phải: l.Add(10) → hết chỗ, List cấp dãy mới 8 ô, chép 4 phần tử sang rồi thêm 10 (Count 5, Capacity 8).",
           "Xong: so sánh hai cột."],
    data="Điểm: 8, 6, 9, 7, 10",
    key="Mảng cố định kích thước khi tạo. List<T> tự lo việc mở rộng, người dùng chỉ cần gọi Add.",
    cs="new List<int>() có Capacity 0; lần Add đầu cấp 4, sau đó nhân đôi (chi tiết cài đặt của .NET, chỉ nêu để minh họa). Ghi ngoài mảng ném IndexOutOfRangeException.")
row(1, L, "Khái niệm và Khởi tạo List<T>", R, N,
    goal="Phân biệt 4 cách khởi tạo List<T> và hiểu Count (số phần tử thật) khác Capacity (chỗ đã chuẩn bị).",
    form="1 cột, 4 tab",
    visual="Dãy ô List với ô dự phòng nét đứt; hai ô biến lớn Count và Capacity. Tab 'Từ mảng' đặt dãy ô mảng và dãy ô List cạnh nhau để thấy là hai bản riêng.",
    comps="H01", tabs="Rỗng | Có sức chứa ban đầu | Khởi tạo kèm giá trị | Từ mảng có sẵn",
    steps=["Tab Rỗng: new List<string>() → Count 0, Capacity 0; Add(\"An\") → Count 1, Capacity 4.",
           "Tab Sức chứa: new List<string>(10) → 10 ô nét đứt, Count vẫn 0. Thử names[0] khi chưa Add → ArgumentOutOfRangeException: Capacity không phải số phần tử.",
           "Tab Kèm giá trị: new List<string> { \"An\", \"Bình\", \"Chi\" } → mô phỏng cho thấy trình biên dịch gọi Add 3 lần.",
           "Tab Từ mảng: int[] arr = {1, 2, 3}; var list = new List<int>(arr) → dữ liệu được chép. Đổi arr[0] = 99, list[0] vẫn là 1.",
           "Bước cuối mỗi tab: ghi chú Count và Capacity."],
    data="Tên: An, Bình, Chi; mảng {1, 2, 3}",
    key="Count là số phần tử đang có; Capacity là chỗ đã cấp sẵn. Chỉ truy cập được chỉ số 0 đến Count − 1.",
    cs="list[i] với i ≥ Count ném ArgumentOutOfRangeException (khác mảng: IndexOutOfRangeException). Collection initializer thực chất gọi Add.")
row(1, L, "Thực hành Khởi tạo List<T>", R, K, note="Trùng ý với 'Khái niệm và Khởi tạo List<T>'. Đính kèm lại mô phỏng đó.")
row(1, L, "Thêm, Truy cập và Xóa phần tử trong List<T>", R, N,
    goal="Thấy các phần tử dịch chuyển thế nào khi Insert/Remove, và lỗi khi dùng chỉ số ngoài phạm vi.",
    form="1 cột, 4 tab",
    visual="Dãy ô có chỉ số. Ô mới trượt vào (vàng); khi chèn ở giữa, các ô phía sau trượt phải và chỉ số đổi (mũi tên 1→2, 2→3). Ô bị xóa mờ dần rồi các ô sau trượt trái.",
    comps="H01, H03", tabs="Add và Insert | Truy cập và sửa | Remove và RemoveAt | Lỗi chỉ số",
    steps=["Add/Insert: fruits = {\"Táo\", \"Cam\"}. Add(\"Xoài\") vào cuối (chỉ số 2). Insert(1, \"Lê\") → \"Cam\", \"Xoài\" dịch phải.",
           "Truy cập: fruits[0] đọc \"Táo\"; fruits[2] = \"Nho\" sửa tại chỗ, Count không đổi.",
           "Remove: Remove(\"Cam\") xóa lần xuất hiện đầu tiên, các ô sau dịch trái, trả true. Remove(\"Dứa\") không có → trả false, không lỗi. RemoveAt(0) xóa theo vị trí.",
           "Lỗi chỉ số: danh sách 3 phần tử, đọc fruits[3] → ArgumentOutOfRangeException; nhắc chỉ số hợp lệ 0..Count − 1.",
           "Bước cuối mỗi tab: tóm tắt trạng thái danh sách."],
    data="Táo, Cam, Xoài, Lê, Nho, Dứa",
    key="Thêm/xóa ở giữa làm dịch chỉ số các phần tử phía sau. Remove theo giá trị trả về bool, không ném lỗi khi không tìm thấy.",
    cs="RemoveAt/indexer với chỉ số sai ném ArgumentOutOfRangeException. Remove chỉ xóa lần xuất hiện đầu tiên.")

L = "Thao tác nâng cao với List<T>"
row(1, L, "Tìm kiếm phần tử trong List<T>", R, N,
    goal="Thấy mọi phép tìm trên List đều quét từ đầu, và phân biệt kết quả của IndexOf, Find, FindAll, Exists.",
    form="1 cột, 4 tab",
    visual="Dãy ô với 'kính lúp' (ô đang xét nền vàng) di chuyển từng ô; dưới mỗi ô hiện kết quả điều kiện true/false; ô khớp viền xanh; kết quả trả về hiện trong thẻ riêng (một số, một phần tử, hay một List mới).",
    comps="H01, H22", tabs="IndexOf / Contains | Find | FindAll | Exists",
    steps=["IndexOf(9): xét 7, 4, 9 → trả 2 và dừng ở lần xuất hiện đầu. IndexOf(5) quét hết → −1. Contains(x) tương đương IndexOf(x) ≥ 0.",
           "Find(x => x > 8): xét từng ô, hiện kết quả điều kiện; trả 9 (phần tử đầu tiên khớp). Find(x => x > 100) → trả 0 (default của int).",
           "FindAll(x => x >= 6): duyệt hết, thả các ô khớp vào một List mới {7, 9, 6, 9}; danh sách gốc giữ nguyên.",
           "Exists(x => x < 4): dừng ngay khi gặp 3 → true.",
           "Bước cuối: ô đếm 'số ô đã xét' cho từng phép tìm."],
    data="scores = {7, 4, 9, 6, 9, 3}",
    key="IndexOf trả vị trí, Find trả phần tử, FindAll trả danh sách mới, Exists trả true/false. Tất cả đều duyệt tuần tự.",
    cs="Find không thấy trả default(T) (0 với int, null với đối tượng) nên dễ nhầm với giá trị hợp lệ; khi cần phân biệt dùng FindIndex hoặc Exists.")
row(1, L, "Thực hành tìm kiếm trong List<T>", R, K, note="Trùng ý 'Tìm kiếm phần tử trong List<T>'. Đính kèm lại mô phỏng đó.")
row(1, L, "Sắp xếp và Đảo ngược List<T>", R, N,
    goal="Thấy kết quả Sort() với số, chuỗi, đối tượng theo tiêu chí riêng, và Reverse() chỉ lật thứ tự.",
    form="1 cột, 4 tab",
    visual="Hai hàng ô 'trước' và 'sau', đường nối từ vị trí cũ tới vị trí mới. Tab đối tượng dùng thẻ sản phẩm (tên + giá); hiện vài phép so sánh mẫu với kết quả âm/0/dương.",
    comps="H01, H17", tabs="Sort số | Sort chuỗi | Sort theo thuộc tính | Reverse",
    steps=["Sort số: {5, 2, 8, 1} → {1, 2, 5, 8}. Nhấn mạnh Sort sửa chính danh sách, không tạo danh sách mới.",
           "Sort chuỗi: {\"Minh\", \"An\", \"Lan\", \"Bao\"} → {\"An\", \"Bao\", \"Lan\", \"Minh\"}.",
           "Sort theo thuộc tính: products.Sort((a, b) => a.Price.CompareTo(b.Price)); hiện 2-3 phép so sánh mẫu (5000 vs 12000 → âm: a đứng trước). Đổi thành b.Price.CompareTo(a.Price) → giảm dần.",
           "Reverse: {1, 2, 5, 8} → {8, 5, 2, 1}. Thử Reverse trên {5, 2, 8, 1} → {1, 8, 2, 5} để thấy Reverse không sắp xếp.",
           "Bước cuối mỗi tab: hai hàng trước/sau."],
    data="Số {5, 2, 8, 1}; tên không dấu; sản phẩm: Bút 5000, Vở 12000, Thước 8000",
    key="Sort() và Reverse() thay đổi trực tiếp danh sách; sắp theo tiêu chí riêng bằng cách truyền hàm so sánh.",
    cs="Sort chuỗi mặc định theo quy tắc văn hóa (culture) nên ví dụ dùng tên không dấu cho dễ đoán. List.Sort không ổn định (các phần tử bằng nhau có thể đổi thứ tự); cần ổn định thì dùng OrderBy.")
row(1, L, "Sắp xếp danh sách", A, K, note="Hoạt động đã có nội dung.")
row(1, L, "Duyệt và Chuyển đổi List<T>", R, N,
    goal="Phân biệt for và foreach, hiểu vì sao không được thêm/xóa phần tử trong foreach, và chuyển đổi List sang dạng khác.",
    form="1 cột, 3 tab",
    visual="Dãy ô với con trỏ duyệt. Tab lỗi có thẻ 'phiên bản danh sách' (version 1 → 2) để giải thích vì sao vòng lặp phát hiện danh sách bị sửa. Tab chuyển đổi hiện danh sách mới xuất hiện bên dưới, danh sách gốc giữ nguyên.",
    comps="H01, H03", tabs="for và foreach | Xóa trong foreach | Chuyển đổi",
    steps=["for: i chạy 0..Count − 1, prices[i] = prices[i] * 2 sửa trực tiếp từng phần tử.",
           "foreach: biến p lần lượt nhận từng giá trị để đọc (tính tổng). Ghi chú: gán p = ... là lỗi biên dịch CS1656.",
           "Xóa trong foreach: nums = {1, 2, 3, 4}; khi gặp 2 gọi nums.Remove(2) → phiên bản danh sách đổi → lần lặp kế tiếp ném InvalidOperationException \"Collection was modified; enumeration operation may not execute.\"",
           "Cách đúng: nums.RemoveAll(n => n % 2 == 0) → {1, 3}.",
           "Chuyển đổi: prices.ConvertAll(p => p * 1.1) tạo List<double> mới; ToArray() tạo mảng; danh sách gốc không đổi."],
    data="prices = {100, 200, 300}; nums = {1, 2, 3, 4}",
    key="Dùng for khi cần chỉ số hoặc sửa phần tử, foreach để đọc. Không thêm/xóa phần tử trong foreach; dùng RemoveAll hoặc for duyệt ngược.",
    cs="InvalidOperationException ném ở lần MoveNext sau khi danh sách bị thay đổi. CS1656: không gán được cho biến lặp foreach.")
row(1, L, "Bài tập: Xây dựng ứng dụng quản lý danh sách sản phẩm", A, K, note="Bài tập nộp, giữ nguyên.")

L = "Giới thiệu về Dictionary<K,V>"
row(1, L, "Giới thiệu Dictionary<K,V> và tầm quan trọng của Key-Value Pair", R, N,
    goal="Hình dung Dictionary như tủ có ngăn dán nhãn: tra theo khóa đi thẳng tới giá trị, không phải quét từng phần tử.",
    form="2 cột song song: List<Student> | Dictionary<string, string>",
    visual="Trái: dãy thẻ sinh viên, kính lúp quét từng thẻ. Phải: tủ 6 ngăn, nhãn là mã SV; mũi tên đi thẳng tới ngăn cần tìm. Dưới mỗi cột là bảng đếm số lần so sánh.",
    comps="H01, H02, H22", tabs="Tìm SV05 | Tìm SV09 (không có)",
    steps=["Cả hai chứa 6 sinh viên SV01..SV06.",
           "Trái: Find(s => s.Id == \"SV05\") xét SV01, SV02... mỗi bước tăng bộ đếm.",
           "Phải: dict[\"SV05\"] → tính vị trí từ khóa và đi thẳng tới ngăn SV05 (1 lần tra).",
           "Kết quả: trái 5 lần so sánh, phải 1 lần tra.",
           "Tab SV09: trái quét hết 6 thẻ rồi trả null; phải TryGetValue trả false ngay."],
    data="SV01 An, SV02 Bình, SV03 Chi, SV04 Dũng, SV05 Hà, SV06 Khoa",
    key="Mỗi khóa là duy nhất và dẫn thẳng tới một giá trị. Tra theo khóa là thế mạnh của Dictionary.",
    cs="Tra theo khóa trung bình O(1) nhờ bảng băm; mô phỏng chỉ minh họa ý 'đi thẳng tới ngăn', không vẽ chi tiết bảng băm.")
row(1, L, "Thực hành Khởi tạo Dictionary<K,V>", A, K, note="Hoạt động đã có nội dung.")
row(1, L, "Thêm, Truy cập và Cập nhật phần tử trong Dictionary<K,V>", R, N,
    goal="Phân biệt Add và gán bằng chỉ mục [key], và nhận ra hai lỗi hay gặp nhất.",
    form="1 cột, 4 tab",
    visual="Tủ ngăn kéo: ngăn mới trượt vào (xanh), ngăn cập nhật nháy vàng, thẻ lỗi đỏ khi Add trùng khóa hoặc đọc khóa không có.",
    comps="H02, H03", tabs="Add | Gán bằng [ ] | Add trùng khóa | Đọc khóa không có",
    steps=["Add: stock.Add(\"thước\", 3) → thêm ngăn mới, Count 3.",
           "Gán: stock[\"vở\"] = 8 → cập nhật ngăn có sẵn; stock[\"tẩy\"] = 2 → khóa chưa có nên thêm mới.",
           "Add trùng: stock.Add(\"bút\", 7) → ArgumentException \"An item with the same key has already been added. Key: bút\".",
           "Đọc khóa không có: int n = stock[\"kéo\"] → KeyNotFoundException \"The given key 'kéo' was not present in the dictionary.\"",
           "Cách an toàn: stock.TryGetValue(\"kéo\", out n) → false, n = 0, không lỗi."],
    data="stock: bút = 10, vở = 5",
    key="Add chỉ để thêm (trùng khóa là lỗi); [key] = value vừa thêm vừa cập nhật; đọc an toàn bằng TryGetValue.",
    cs="Thông điệp lỗi theo .NET hiện đại (.NET Core/.NET 5+).")
row(1, L, "Bài tập: Thao tác cơ bản với Dictionary<K,V>", A, K, note="Bài tập nộp, giữ nguyên.")
row(1, L, "Kiểm tra sự tồn tại và Xóa phần tử trong Dictionary<K,V>", R, T,
    goal="Dùng đúng ContainsKey, ContainsValue và Remove; thấy kiểm tra giá trị chậm hơn kiểm tra khóa.",
    form="1 cột, 3 tab",
    visual="Tủ ngăn kéo: ContainsKey đi thẳng tới ngăn; ContainsValue mở lần lượt từng ngăn; ngăn bị Remove mờ đỏ rồi biến mất.",
    comps="H02", tabs="ContainsKey | ContainsValue | Remove",
    steps=["ContainsKey(\"vở\") → true sau 1 lần tra; ContainsKey(\"kéo\") → false.",
           "ContainsValue(5) → mở lần lượt từng ngăn đến khi gặp giá trị 5 (đếm số ngăn đã mở).",
           "Remove(\"bút\") → true, Count giảm; Remove(\"kéo\") → false, không lỗi.",
           "Remove(\"vở\", out int old) → lấy luôn giá trị vừa xóa (old = 5)."],
    data="stock: bút = 10, vở = 5, thước = 3",
    key="Kiểm tra khóa nhanh, kiểm tra giá trị phải duyệt; Remove khóa không có chỉ trả false.",
    cs="Remove(key, out value) có từ .NET Core 2.0.")

L = "Thao tác nâng cao với Dictionary<K,V> và lựa chọn Collections"
row(1, L, "Duyệt qua Dictionary<K,V>", R, T,
    goal="Thấy mỗi lần lặp foreach nhận một cặp KeyValuePair, và cách duyệt riêng Keys hoặc Values.",
    form="1 cột, 3 tab",
    visual="Tủ ngăn kéo với con trỏ duyệt; thẻ 'kv' hiện Key và Value của lần lặp hiện tại; console in dần.",
    comps="H02", tabs="foreach KeyValuePair | Keys | Values",
    steps=["foreach (var kv in scores): lần 1 kv.Key = \"An\", kv.Value = 8 ... in từng dòng.",
           "Keys: duyệt chỉ các khóa, in danh sách tên.",
           "Values: duyệt chỉ giá trị, tính tổng và điểm trung bình (tính từ dữ liệu).",
           "Bước cuối: ghi chú thứ tự duyệt."],
    data="scores: An 8, Bình 6, Chi 9",
    key="foreach trên Dictionary trả về từng cặp khóa – giá trị. Thứ tự duyệt không được bảo đảm.",
    cs="Thứ tự duyệt thường trùng thứ tự thêm nhưng tài liệu không bảo đảm; cần thứ tự thì dùng SortedDictionary hoặc OrderBy.")
row(1, L, "Thực hành Duyệt Dictionary", A, K, note="Hoạt động đã có nội dung.")
row(1, L, "So sánh List<T> và Dictionary<K,V>: Khi nào sử dụng cái nào?", R, N,
    goal="Chọn đúng cấu trúc theo nhu cầu: tra theo khóa, lấy theo vị trí, hay cần giữ giá trị trùng.",
    form="2 cột song song, 3 tab",
    visual="Trái dãy ô List, phải tủ ngăn Dictionary. Dưới mỗi cột là bảng điểm số bước; bước cuối gắn cờ 'phù hợp hơn' cho cột thắng ở tình huống đó.",
    comps="H01, H02, H22", tabs="Tra theo mã | Lấy phần tử thứ 3 | Lưu giá trị trùng",
    steps=["Tra theo mã (8 phần tử, tìm phần tử thứ 6): List quét 6 bước, Dictionary 1 bước → Dictionary phù hợp.",
           "Lấy phần tử thứ 3: List dùng list[2] trong 1 bước; Dictionary không có chỉ số vị trí, phải duyệt → List phù hợp.",
           "Lưu giá trị trùng: List chứa \"An\", \"An\" bình thường; Dictionary Add khóa \"An\" lần 2 → ArgumentException → khóa phải duy nhất.",
           "Bước cuối: bảng tổng hợp 3 tình huống."],
    data="8 sản phẩm SP01..SP08",
    key="List: thứ tự và vị trí. Dictionary: tra nhanh theo khóa duy nhất. Có thể kết hợp, ví dụ Dictionary<string, List<T>>.",
    cs="Số bước được đếm từ dữ liệu thật trong mô phỏng.")
row(1, L, "Tóm tắt và Mẹo sử dụng Collections hiệu quả", R, K, note="Bài tóm tắt.")

# ======================= CHƯƠNG 2: EXCEPTION CƠ BẢN =======================
L = "Giới thiệu về Exception và Tầm quan trọng"
row(2, L, "Exception là gì?", R, N,
    goal="Thấy exception là một đối tượng được 'ném' ra khi có lỗi lúc chạy, và nếu không ai bắt thì chương trình dừng ngay tại đó.",
    form="1 cột, có ô nhập chuỗi (chips gợi ý)",
    visual="Mã nguồn tô dòng đang chạy. Khi lỗi, thẻ Exception đỏ bật lên cạnh dòng gây lỗi; các dòng phía sau mờ đi với nhãn 'không được chạy'; console in thông báo lỗi.",
    comps="H03", tabs="1 tab, ô nhập 'input' (mặc định abc; gợi ý: 20, abc, rỗng, -5)",
    steps=["Console.Write(\"Nhập tuổi: \"); input = giá trị trong ô nhập.",
           "int age = int.Parse(input) → với \"abc\", .NET tạo một đối tượng FormatException và ném ra.",
           "Thẻ đỏ: FormatException, \"The input string 'abc' was not in a correct format.\"",
           "Không có try-catch nên chương trình dừng; hai dòng in phía sau mờ đi.",
           "Với \"20\": chạy hết và in \"Năm sau bạn 21 tuổi\". Với \"-5\": không lỗi (chỉ là dữ liệu sai nghiệp vụ, không phải exception)."],
    data="Ô nhập: abc (mặc định), 20, rỗng, -5",
    key="Exception là sự cố lúc chạy, được biểu diễn bằng một đối tượng; không bắt thì chương trình dừng. Lỗi định dạng khác lỗi nghiệp vụ.",
    cs="Thông điệp theo .NET 8: \"The input string 'abc' was not in a correct format.\" (bản cũ: \"Input string was not in a correct format.\"). Ô nhập cần cho phép chuỗi rỗng và dấu '-'.")
row(2, L, "Tại sao cần xử lý lỗi?", R, T,
    goal="So sánh trải nghiệm người dùng khi chương trình không xử lý lỗi và khi có xử lý lỗi.",
    form="2 cột song song: Không xử lý | Có xử lý",
    visual="Dưới mỗi cột là một 'cửa sổ console' giả lập: trái hiện vết lỗi dài màu đỏ, phải hiện câu thông báo thân thiện. Biểu tượng 'Đơn đã lưu: 1, 2, 3' với ✓/✗.",
    comps="H03, H09", tabs="1 tab: Nhập 3 đơn hàng, đơn 2 nhập sai số lượng",
    steps=["Cả hai xử lý 3 đơn hàng: số lượng \"2\", \"hai\", \"5\".",
           "Đơn 1: cả hai xử lý và lưu bình thường.",
           "Đơn 2: trái ném FormatException, chương trình sập; phải bắt lỗi, báo \"Số lượng không hợp lệ, bỏ qua đơn 2\".",
           "Đơn 3: trái không bao giờ được xử lý; phải xử lý bình thường.",
           "Kết quả: trái lưu 1/3 đơn và hiện thông báo khó hiểu; phải lưu 2/3 đơn và báo rõ đơn lỗi."],
    data="Số lượng: 2, hai, 5",
    key="Xử lý lỗi giúp chương trình không sập, giữ được phần dữ liệu tốt và báo lỗi dễ hiểu.",
    cs="Có thể đính kèm lại ở 'Giới thiệu về khối try-catch' nếu không làm riêng.")
row(2, L, "Các loại Exception phổ biến trong C#", R, N,
    goal="Nhận diện các exception hay gặp qua chính dòng code gây ra chúng.",
    form="1 cột, 6 tab (bộ sưu tập thẻ lỗi)",
    visual="Mỗi tab là một 'thẻ bài' exception: biểu tượng, tên loại, dòng code gây lỗi, thông điệp thật của .NET, một dòng 'Cách phòng tránh'. Khung bộ nhớ hiện giá trị biến dẫn tới lỗi.",
    comps="H03", tabs="FormatException | DivideByZeroException | IndexOutOfRangeException | NullReferenceException | InvalidCastException | OverflowException",
    steps=["Mỗi tab 3 bước: (1) chuẩn bị dữ liệu, (2) chạy dòng gây lỗi, (3) thẻ lỗi + cách phòng tránh.",
           "Format: int.Parse(\"12a\") → \"The input string '12a' was not in a correct format.\" → dùng int.TryParse.",
           "DivideByZero: int a = 10, b = 0; a / b → \"Attempted to divide by zero.\" → kiểm tra b != 0. Ghi chú: double chia 0 không lỗi mà ra ∞.",
           "IndexOutOfRange: int[] arr = new int[3]; arr[3] → \"Index was outside the bounds of the array.\" → kiểm tra i < arr.Length.",
           "NullReference: string s = null; s.Length → \"Object reference not set to an instance of an object.\" → kiểm tra null hoặc dùng s?.Length.",
           "InvalidCast: object o = \"abc\"; (int)o → \"Unable to cast object of type 'System.String' to type 'System.Int32'.\" → dùng is/as.",
           "Overflow: checked { int x = int.MaxValue; x = x + 1; } → \"Arithmetic operation resulted in an overflow.\" Ghi chú: không có checked thì không lỗi mà quay vòng thành số âm."],
    data="Dữ liệu cố định từng tab",
    key="Tên loại exception cho biết nguyên nhân; đọc tên loại là bước đầu tiên để sửa lỗi.",
    cs="Thông điệp theo .NET hiện đại. OverflowException chỉ xảy ra trong ngữ cảnh checked (hoặc khi Parse/Convert số quá lớn).")
row(2, L, "Fix chương trình gặp lỗi runtime", A, K, note="Hoạt động đã có nội dung.")

L = "Cấu trúc try-catch cơ bản"
row(2, L, "Giới thiệu về khối try-catch", R, N,
    goal="Thấy đường đi của chương trình qua try-catch trong hai trường hợp: không lỗi và có lỗi.",
    form="1 cột, 2 tab",
    visual="Sơ đồ luồng nhỏ đặt cạnh mã: khối try, khối catch, khối 'sau try-catch'. Đường đi tô xanh khi bình thường, tô đỏ khi nhảy sang catch. Các dòng còn lại trong try bị bỏ qua hiện mờ.",
    comps="H04, H03", tabs="Không lỗi (\"25\") | Có lỗi (\"abc\")",
    steps=["Không lỗi: vào try → Parse thành công → in \"Bạn 25 tuổi\" → hết try → catch bị bỏ qua → in \"Kết thúc\".",
           "Có lỗi: vào try → Parse ném FormatException.",
           "Dòng in tuổi còn lại trong try bị bỏ qua (mờ), chương trình nhảy thẳng vào catch.",
           "catch in \"Vui lòng nhập số\".",
           "Chương trình chạy tiếp sau khối, in \"Kết thúc\"."],
    data="\"25\" và \"abc\"",
    key="Lỗi trong try làm phần còn lại của try bị bỏ qua; catch chỉ chạy khi có lỗi; chương trình tiếp tục sau khối try-catch.")
row(2, L, "Thực hành với try-catch đơn giản", A, K, note="Hoạt động đã có nội dung.")
row(2, L, "Đối tượng Exception và các thuộc tính", R, T,
    goal="Biết lấy thông tin từ biến ex trong catch: Message cho người dùng, StackTrace cho lập trình viên.",
    form="1 cột",
    visual="Bảng kiểm tra đối tượng (inspector) giống cửa sổ Watch: hiện lần lượt GetType().Name, Message, StackTrace (rút gọn 2 dòng), InnerException = null.",
    comps="H03, H07", tabs="1 tab",
    steps=["Main gọi Chia(10, 0); trong Chia, 10 / 0 ném DivideByZeroException.",
           "catch (DivideByZeroException ex) bắt được; ex là một đối tượng.",
           "Hiện ex.GetType().Name = \"DivideByZeroException\".",
           "Hiện ex.Message = \"Attempted to divide by zero.\"",
           "Hiện ex.StackTrace: \"at Program.Chia(Int32 a, Int32 b) in Program.cs:line 12\" và \"at Program.Main() in Program.cs:line 5\".",
           "Xong: Message dùng để báo người dùng, StackTrace dùng để tìm chỗ lỗi."],
    data="Chia(10, 0)",
    key="Exception là đối tượng có thuộc tính; StackTrace cho biết lỗi đi qua những hàm nào.",
    cs="Số dòng trong StackTrace chỉ có khi có tệp .pdb (thường là bản Debug).")
row(2, L, "Hiển thị thông tin Exception chi tiết", A, K, note="Hoạt động đã có nội dung.")
row(2, L, "Bài tập: Xây dựng máy tính chia số an toàn", A, K, note="Bài tập nộp, giữ nguyên.")

L = "Bắt và xử lý nhiều loại Exception"
row(2, L, "Bắt nhiều loại Exception trong một khối try", R, N,
    goal="Thấy C# so khớp exception với từng catch từ trên xuống và chỉ chạy catch đầu tiên phù hợp.",
    form="1 cột, 4 tab theo dữ liệu nhập",
    visual="Các catch hiển thị như những cửa lọc xếp dọc. Exception (viên đỏ) rơi từ trên xuống: cửa không khớp hiện ✗ xám, cửa khớp hiện ✓ xanh. Tab không lỗi: viên đỏ không xuất hiện, các cửa đều xám.",
    comps="H05, H03", tabs="\"abc\" | \"0\" | \"99999999999\" | \"5\"",
    steps=["Code: int n = int.Parse(input); int r = 100 / n; ba catch: FormatException, DivideByZeroException, OverflowException.",
           "\"abc\": Parse ném FormatException → khớp cửa 1.",
           "\"0\": Parse được 0; 100 / 0 ném DivideByZeroException → cửa 1 ✗, cửa 2 ✓.",
           "\"99999999999\": số vượt int → Parse ném OverflowException → cửa 1 ✗, cửa 2 ✗, cửa 3 ✓.",
           "\"5\": không lỗi, r = 20, không catch nào chạy.",
           "Bước cuối mỗi tab: chỉ đúng một catch (hoặc không catch nào) được chạy."],
    data="abc, 0, 99999999999, 5",
    key="Mỗi catch xử lý một loại lỗi. C# xét từ trên xuống và dừng ở catch khớp đầu tiên.",
    cs="int.Parse với số vượt phạm vi int ném OverflowException (không phải FormatException).")
row(2, L, "Thực hành bắt nhiều Exception", A, K, note="Hoạt động đã có nội dung.")
row(2, L, "Thứ tự bắt Exception và Exception chung", R, N,
    goal="Hiểu vì sao catch cụ thể phải đứng trước catch (Exception), dựa trên quan hệ cha – con giữa các lớp exception.",
    form="1 cột, 3 tab",
    visual="Cây phân cấp: Exception → SystemException → FormatException, NullReferenceException; SystemException → ArithmeticException → DivideByZeroException, OverflowException. Khi so khớp, tô đường đi từ loại exception lên các lớp cha. Kết hợp bộ lọc catch xếp dọc.",
    comps="H06, H05, H03", tabs="Thứ tự đúng | Exception chung bắt lỗi lạ | Thứ tự sai",
    steps=["Thứ tự đúng: catch Format, catch DivideByZero, catch Exception. Ném DivideByZeroException → khớp catch 2.",
           "Exception chung: ném NullReferenceException → không khớp 2 catch đầu; leo cây: NullReferenceException là một Exception → catch (Exception) bắt được.",
           "Thứ tự sai: catch (Exception) đứng đầu → trình biên dịch báo lỗi CS0160 tại catch phía sau: \"A previous catch clause already catches all exceptions of this or of a super type ('Exception')\".",
           "Bước cuối: catch (Exception) là lưới an toàn cuối cùng, đặt sau mọi catch cụ thể."],
    data="Các loại exception cố định",
    key="Catch cụ thể trước, chung sau. Một catch bắt được mọi exception là lớp con của kiểu nó khai báo.",
    cs="DivideByZeroException và OverflowException đều kế thừa ArithmeticException. Thứ tự sai là lỗi biên dịch, không phải lỗi lúc chạy.")
row(2, L, "Bài tập: Kiểm tra thứ tự bắt Exception", A, K, note="Bài tập nộp, giữ nguyên.")

L = "Khối finally và các trường hợp đặc biệt"
row(2, L, "Khối finally là gì và khi nào sử dụng", R, N,
    goal="Thấy finally luôn chạy: khi không lỗi, khi có lỗi được bắt, và cả khi try có return.",
    form="1 cột, 3 tab",
    visual="Sơ đồ luồng try → catch → finally → sau khối. Đường đi tô theo từng tình huống; khối finally luôn được tô kèm nhãn 'luôn chạy'. Console in theo đúng thứ tự thực thi.",
    comps="H04", tabs="Không lỗi | Có lỗi được bắt | Có return trong try",
    steps=["Không lỗi: try chạy hết → bỏ qua catch → finally in \"Dọn dẹp\" → \"Sau khối\".",
           "Có lỗi: try ném lỗi → catch in thông báo → finally in \"Dọn dẹp\" → \"Sau khối\".",
           "Có return: hàm return 10 trong try → finally vẫn chạy và in \"Dọn dẹp\" trước khi hàm thật sự trả về; dòng \"Sau khối\" trong hàm không chạy.",
           "Bước cuối: so sánh thứ tự in ở 3 tab."],
    data="Code cố định",
    key="Đặt mã dọn dẹp (đóng tệp, đóng kết nối, ghi log) vào finally.",
    cs="finally không chạy nếu tiến trình bị kết thúc đột ngột (Environment.FailFast, tắt nguồn).")
row(2, L, "Sử dụng finally để đảm bảo tài nguyên được giải phóng", R, N,
    goal="Thấy finally giúp đóng tệp ngay cả khi lỗi xảy ra giữa chừng.",
    form="2 cột song song: Không finally | Có finally",
    visual="Biểu tượng tệp data.txt có đèn trạng thái (Đóng / Đang mở / Vẫn mở – rò rỉ). Sơ đồ luồng thu nhỏ dưới mỗi cột.",
    comps="H08, H04, H03", tabs="Lỗi khi đang đọc",
    steps=["Hai cột mở data.txt bằng StreamReader → đèn 'Đang mở'.",
           "Đọc dòng 1 bình thường.",
           "Dòng 2 không phải số, int.Parse ném FormatException.",
           "Trái: Close() nằm sau dòng lỗi nên không chạy; catch báo lỗi; đèn vẫn 'Đang mở' (đỏ).",
           "Phải: catch báo lỗi, finally gọi reader.Close() → đèn 'Đóng'.",
           "Xong: so sánh trạng thái tệp hai bên."],
    data="data.txt: 10 / abc / 30",
    key="Đưa lệnh giải phóng tài nguyên vào finally để nó chạy trong mọi trường hợp.",
    cs="Cách gọn hơn là using (học ở chương sau).")
row(2, L, "Tóm tắt về try-catch-finally", R, K, note="Bài tóm tắt.")

# ======================= CHƯƠNG 3: EXCEPTION NÂNG CAO =======================
L = "Xử lý Đa Ngoại lệ và Thứ tự Catch"
row(3, L, "Giới thiệu về Xử lý Đa Ngoại lệ", R, K, note="Trùng ý 'Bắt nhiều loại Exception trong một khối try' (chương 2). Đính kèm lại mô phỏng đó.")
row(3, L, "Cấu trúc try-catch với nhiều catch block", R, K, note="Trùng ý chương 2. Đính kèm lại 'Bắt nhiều loại Exception trong một khối try'.")
row(3, L, "Minh họa thứ tự catch block", R, K, note="Trùng ý 'Thứ tự bắt Exception và Exception chung' (chương 2). Đính kèm lại mô phỏng đó.")
row(3, L, "Phân tích ví dụ try-catch phức tạp", A, T,
    goal="Thấy exception 'nổi lên' qua các hàm gọi nhau cho tới khi gặp catch phù hợp gần nhất.",
    form="1 cột, 2 tab",
    visual="Ngăn xếp lời gọi: các khung Main (đáy), ProcessOrder, ParseQuantity (đỉnh). Exception là quả bóng đỏ đi lên từng khung; khung không có catch phù hợp bị đóng lại (mờ, các dòng còn lại không chạy); khung có catch bắt lấy (viền xanh).",
    comps="H07, H03", tabs="Bắt ở Main | Bắt ngay trong ProcessOrder",
    steps=["Main gọi ProcessOrder(\"hai\") → khung mới chồng lên.",
           "ProcessOrder gọi ParseQuantity(\"hai\") → khung thứ ba.",
           "int.Parse ném FormatException tại ParseQuantity; khung này không có catch → đóng lại.",
           "Bóng đỏ lên ProcessOrder: tab 1 không có catch → khung đóng, dòng tính tiền phía sau không chạy; tab 2 có catch → bắt tại đây, trả về giá trị mặc định.",
           "Tab 1: bóng lên Main, catch ở Main bắt và in thông báo.",
           "Xong: so sánh hai nơi bắt lỗi."],
    data="Số lượng \"hai\"",
    key="Exception đi ngược ngăn xếp lời gọi tới catch gần nhất; mọi dòng còn lại của các hàm trung gian bị bỏ qua.",
    note="Học liệu hiện là hoạt động phân tích; mô phỏng đính kèm bổ trợ.")
row(3, L, "Sửa lỗi trong đoạn code có nhiều ngoại lệ", A, K, note="Hoạt động đã có nội dung.")

L = "Khối finally và Đảm bảo Giải phóng Tài nguyên"
row(3, L, "Vai trò của khối finally", R, K, note="Trùng ý 'Khối finally là gì và khi nào sử dụng' (chương 2). Đính kèm lại mô phỏng đó.")
row(3, L, "finally trong thực tế: Đóng kết nối, giải phóng tài nguyên", R, T,
    goal="Thấy hậu quả tích lũy khi quên giải phóng tài nguyên: số kết nối bị bỏ ngỏ tăng dần sau mỗi lần lỗi.",
    form="2 cột song song: Không finally | Có finally",
    visual="Mỗi cột có 'bể kết nối' gồm 5 ô; ô đang dùng màu xanh, ô bị bỏ ngỏ màu đỏ. Bộ đếm 'Kết nối đang mở' lớn ở trên.",
    comps="H08, H22", tabs="Gọi hàm 3 lần, cả 3 lần đều lỗi",
    steps=["Mỗi lần gọi LoadData(): mở kết nối → truy vấn → đóng kết nối.",
           "Lần 1: truy vấn lỗi. Trái: Close bị bỏ qua, đếm = 1. Phải: finally đóng, đếm = 0.",
           "Lần 2: trái đếm = 2; phải vẫn 0.",
           "Lần 3: trái đếm = 3; phải vẫn 0.",
           "Thêm bước: lần 6 bên trái vượt 5 ô → báo 'Hết kết nối' (minh họa giới hạn bể kết nối).",
           "Xong: so sánh bộ đếm."],
    data="Bể kết nối giả lập 5 ô",
    key="Rò rỉ tài nguyên không gây lỗi ngay mà tích lũy dần cho tới khi hệ thống cạn tài nguyên.",
    cs="Kết nối là giả lập (không dùng thư viện cơ sở dữ liệu thật); con số giới hạn bể chỉ để minh họa.")
row(3, L, "Sử dụng finally để đóng file", R, K, note="Trùng ý 'Sử dụng finally để đảm bảo tài nguyên được giải phóng' (chương 2). Đính kèm lại.")
row(3, L, "Quan sát hành vi của finally với và không có ngoại lệ", A, K, note="Hoạt động đã có nội dung.")
row(3, L, "using statement như một cách thay thế cho finally (cho IDisposable)", R, N,
    goal="Thấy using tự gọi Dispose khi rời khối (kể cả khi có lỗi), tương đương try/finally viết tay.",
    form="2 cột song song: try/finally viết tay | using",
    visual="Biểu tượng tài nguyên với trạng thái Đang mở / Đã Dispose. Mũi tên chú thích 'trình biên dịch dịch using thành try/finally' nối hai cột. Đường đi tô theo tình huống.",
    comps="H08, H04", tabs="Không lỗi | Có lỗi giữa chừng | Khai báo using (C# 8)",
    steps=["Không lỗi: hai cột mở StreamReader, đọc, rời khối → cả hai gọi Dispose (trái do finally, phải do using).",
           "Có lỗi: lỗi khi đọc → trái nhảy vào finally gọi Dispose; phải cũng Dispose tự động trước khi lỗi lan ra ngoài.",
           "Khai báo using: using var reader = new StreamReader(...); Dispose chạy khi hết phạm vi (cuối hàm).",
           "Bước cuối: hai cột cùng kết quả, cột phải ngắn gọn hơn."],
    data="data.txt",
    key="Với đối tượng IDisposable (tệp, kết nối...), ưu tiên dùng using thay cho try/finally viết tay.",
    cs="Khai báo 'using var' có từ C# 8; đối tượng Dispose khi ra khỏi phạm vi chứa nó.")
row(3, L, "Bài tập: Quản lý tài nguyên hệ thống an toàn với try-catch-finally", A, K, note="Bài tập nộp, giữ nguyên.")

L = "Tạo và Sử dụng Ngoại lệ Tùy chỉnh (Custom Exceptions)"
row(3, L, "Tại sao cần ngoại lệ tùy chỉnh?", R, T,
    goal="So sánh bắt lỗi bằng cách dò chuỗi thông điệp với bắt theo một kiểu exception riêng.",
    form="2 cột song song: Exception chung | InsufficientFundsException",
    visual="Thẻ exception hai bên: trái chỉ có Message; phải có thêm Balance và Amount. Một bước 'đổi câu chữ thông điệp' làm điều kiện bên trái sai (thẻ catch chuyển đỏ).",
    comps="H03", tabs="1 tab",
    steps=["Rút 500 khi số dư 200.",
           "Trái: throw new Exception(\"Số dư không đủ\"); catch kiểm tra ex.Message.Contains(\"không đủ\").",
           "Phải: throw new InsufficientFundsException(200, 500); catch (InsufficientFundsException ex).",
           "Có người sửa thông điệp thành \"Tài khoản thiếu tiền\" → điều kiện Contains bên trái sai, lỗi bị xử lý nhầm nhánh.",
           "Phải không bị ảnh hưởng, còn tính được ex.Amount − ex.Balance = 300 để báo thiếu bao nhiêu.",
           "Xong."],
    data="Số dư 200, rút 500",
    key="Exception tùy chỉnh mang ý nghĩa nghiệp vụ và dữ liệu kèm theo, bắt theo kiểu thay vì dò chữ.")
row(3, L, "Hướng dẫn tạo CustomException", R, N,
    goal="Thấy cấu trúc tối thiểu của một exception tự tạo và đường đi throw → catch của nó.",
    form="1 cột",
    visual="Cây kế thừa: nút InsufficientFundsException được vẽ thêm dưới Exception. Thẻ đối tượng exception hiện Message, Balance, Amount. Đường đi throw (từ Withdraw) → catch (ở Main) tô màu.",
    comps="H06, H03, H07", tabs="1 tab",
    steps=["Khai báo class InsufficientFundsException : Exception → cây thêm nút mới.",
           "Constructor nhận balance, amount và gọi : base(\"Số dư không đủ\") → Message được đặt ở lớp cha.",
           "Withdraw(500) khi Balance = 200 → điều kiện amount > Balance đúng → throw new InsufficientFundsException(200, 500).",
           "Thẻ đối tượng: Message, Balance = 200, Amount = 500.",
           "Ở Main, catch (InsufficientFundsException ex) khớp → in \"Thiếu 300\".",
           "Xong."],
    data="Số dư 200, rút 500",
    key="Exception tùy chỉnh là một lớp kế thừa Exception, có thể thêm thuộc tính riêng; tên kết thúc bằng 'Exception'.",
    cs="Khuyến nghị có thêm các constructor chuẩn (không tham số; message; message + innerException); mô phỏng chỉ dùng một constructor cho gọn.")
row(3, L, "Thực hành: Triển khai một CustomException đơn giản", A, K, note="Hoạt động đã có nội dung.")
row(3, L, "Phân tích ví dụ CustomException trong ứng dụng thực tế", A, K, note="Hoạt động đã có nội dung.")

L = "Kiểm tra và Xác thực Dữ liệu Đầu vào"
row(3, L, "Tầm quan trọng của xác thực dữ liệu đầu vào", R, T,
    goal="Thấy dữ liệu sai lọt vào hệ thống gây hậu quả gì khi không được kiểm tra.",
    form="2 cột song song: Không xác thực | Có xác thực",
    visual="Giữa ô nhập và 'kho dữ liệu' (bảng người dùng) có một cổng kiểm soát. Trái không có cổng: bản ghi tuổi −5, email rỗng lọt vào bảng (tô đỏ). Phải: cổng chặn lại kèm lý do.",
    comps="H09, H21", tabs="1 tab: 3 lượt đăng ký",
    steps=["Đăng ký 1: An, 20, an@x.com → cả hai lưu.",
           "Đăng ký 2: Bình, −5, (email rỗng) → trái lưu bản ghi sai; phải chặn: \"Tuổi phải từ 1 đến 120; email không được rỗng\".",
           "Đăng ký 3: Chi, 30, chi@x.com → cả hai lưu.",
           "Tính tuổi trung bình: trái (20 − 5 + 30) / 3 = 15 (sai lệch); phải (20 + 30) / 2 = 25.",
           "Xong."],
    data="(An, 20, an@x.com), (Bình, −5, rỗng), (Chi, 30, chi@x.com)",
    key="Kiểm tra ngay ở đầu vào rẻ hơn nhiều so với sửa dữ liệu hỏng về sau.")
row(3, L, "Các kỹ thuật xác thực dữ liệu: Kiểm tra null, rỗng, định dạng, phạm vi", R, N,
    goal="Thấy một giá trị đi qua lần lượt 4 lớp kiểm tra và dừng ở lớp đầu tiên không đạt.",
    form="1 cột, 5 tab theo giá trị đầu vào",
    visual="Dải 4 đèn kiểm tra nằm ngang: Có giá trị? → Không chỉ khoảng trắng? → Là số? (int.TryParse) → Trong khoảng 1–120? Đèn sáng xanh khi đạt, đỏ khi trượt, các đèn phía sau tắt. Thông báo lỗi tương ứng hiện dưới đèn đỏ.",
    comps="H09", tabs="null | \"   \" | \"abc\" | \"150\" | \"25\"",
    steps=["Mỗi tab đi qua các đèn theo thứ tự, mỗi đèn một bước.",
           "null → trượt đèn 1: \"Vui lòng nhập tuổi\".",
           "\"   \" → qua đèn 1, trượt đèn 2.",
           "\"abc\" → qua đèn 1, 2; TryParse trả false → trượt đèn 3: \"Tuổi phải là số\".",
           "\"150\" → qua 3 đèn đầu, trượt đèn 4: \"Tuổi phải từ 1 đến 120\".",
           "\"25\" → 4 đèn xanh: hợp lệ."],
    data="null, \"   \", \"abc\", \"150\", \"25\"",
    key="Kiểm tra theo thứ tự từ chung đến cụ thể, mỗi trường hợp một thông báo rõ ràng; dùng TryParse thay vì Parse + try-catch cho dữ liệu người dùng.",
    cs="string.IsNullOrWhiteSpace gộp được 2 kiểm tra đầu; mô phỏng tách riêng để dễ quan sát.")
row(3, L, "Thực hành: Viết hàm xác thực số nguyên", A, K, note="Hoạt động đã có nội dung.")
row(3, L, "Bài tập: Xây dựng Bộ xác thực Thông tin Người dùng (User Registration Validator)", A, K, note="Bài tập nộp, giữ nguyên.")
row(3, L, "Các thư viện hỗ trợ xác thực dữ liệu", R, K, note="Bài tổng quan thư viện (DataAnnotations, FluentValidation...); không có quá trình chạy cần mô phỏng.")
row(3, L, "Thử nghiệm các trường hợp dữ liệu đầu vào hợp lệ/không hợp lệ", A, K, note="Hoạt động đã có nội dung; có thể đính kèm thêm mô phỏng 'Các kỹ thuật xác thực dữ liệu'.")

L = "Thực hành Tổng hợp và Best Practices"
row(3, L, "Các nguyên tắc vàng trong xử lý ngoại lệ", R, T,
    goal="Nhận ra 3 lỗi thực hành phổ biến và cách sửa, qua so sánh 'nên tránh' và 'nên làm'.",
    form="2 cột song song (Nên tránh | Nên làm), 3 tab",
    visual="Cột trái viền đỏ, cột phải viền xanh. Hộp 'Thông tin lỗi còn lại' ở cuối mỗi cột cho thấy thông tin còn hay mất (stack trace đủ dòng hay bị cắt).",
    comps="H03, H07, H22", tabs="Nuốt lỗi | throw ex và throw | Dùng exception cho luồng thường",
    steps=["Nuốt lỗi: trái catch { } rỗng → lỗi biến mất, tổng tiền ra 0 mà không ai biết; phải ghi log rồi báo lỗi.",
           "throw ex và throw: lỗi gốc ở ParseQuantity; trái 'throw ex;' → stack trace bắt đầu từ hàm có catch (mất dòng gốc); phải 'throw;' → giữ dòng ParseQuantity.",
           "Luồng thường: đọc 1.000 dòng, 300 dòng sai. Trái int.Parse trong try-catch → 300 exception được ném; phải int.TryParse → 0 exception. Ô đếm 'Số exception đã ném'.",
           "Bước cuối mỗi tab: nguyên tắc rút ra."],
    data="Dữ liệu cố định từng tab",
    key="Không nuốt lỗi; ném lại bằng throw; để giữ stack trace; dùng TryParse cho dữ liệu có thể sai thay vì dựa vào exception.",
    cs="'throw ex;' đặt lại stack trace về vị trí ném lại; 'throw;' giữ nguyên. Không nêu số mili giây cụ thể về hiệu năng (phụ thuộc máy).")
row(3, L, "Xây dựng một module xử lý dữ liệu với xác thực và try-catch-finally", R, T,
    goal="Thấy xác thực, try-catch và finally phối hợp trong một luồng xử lý hoàn chỉnh.",
    form="1 cột, 3 tab",
    visual="Băng chuyền 4 trạm: Đọc → Xác thực → Xử lý → Ghi log (finally). Mỗi bản ghi là một 'kiện hàng' đi qua; trạm báo lỗi đổi màu; trạm finally luôn sáng.",
    comps="H09, H04, H21", tabs="Dữ liệu hợp lệ | Dữ liệu sai | Lỗi hệ thống",
    steps=["Hợp lệ: đơn (Bút, 2, 5000) qua đủ 4 trạm, tổng 10000, log \"OK\".",
           "Dữ liệu sai: số lượng −1 bị trạm Xác thực chặn (không ném exception), log \"Bỏ qua: số lượng không hợp lệ\".",
           "Lỗi hệ thống: trạm Xử lý gặp exception bất ngờ (ví dụ không ghi được tệp) → catch báo lỗi → finally vẫn ghi log.",
           "Bước cuối mỗi tab: đèn trạm finally luôn sáng."],
    data="Đơn hàng mẫu",
    key="Xác thực xử lý lỗi dự đoán được; try-catch xử lý lỗi bất ngờ; finally dọn dẹp và ghi log trong mọi trường hợp.")
row(3, L, "Đánh giá code mẫu về xử lý lỗi và xác thực", A, K, note="Hoạt động đã có nội dung.")
row(3, L, "Bài tập: Xây dựng Module Xử lý Đơn hàng An toàn", A, K, note="Bài tập nộp, giữ nguyên.")

# ======================= CHƯƠNG 4: FILE VĂN BẢN =======================
L = "Giới thiệu chung về Thao tác File và Stream"
row(4, L, "Vai trò của File I/O trong ứng dụng", R, N,
    goal="Thấy dữ liệu trong biến mất khi chương trình tắt, còn dữ liệu trong tệp thì vẫn còn.",
    form="2 cột song song: Chỉ lưu trong biến | Lưu thêm vào tệp",
    visual="Mỗi cột có 2 khung: 'Bộ nhớ chương trình (RAM)' và 'Ổ đĩa' với biểu tượng notes.txt. Nút trạng thái chương trình: Đang chạy → Đã tắt (RAM xóa trắng có hiệu ứng mờ dần) → Chạy lại.",
    comps="H10, H12", tabs="1 tab",
    steps=["Hai cột thêm 3 ghi chú vào List<string> notes (RAM hiện 3 ô).",
           "Phải gọi File.WriteAllLines(\"notes.txt\", notes) → ổ đĩa xuất hiện tệp 3 dòng.",
           "Tắt chương trình: RAM của cả hai cột bị xóa trắng.",
           "Chạy lại: trái notes rỗng (Count 0).",
           "Phải gọi File.ReadAllLines(\"notes.txt\") → 3 ghi chú trở lại.",
           "Xong."],
    data="Ghi chú: Mua sữa, Học C#, Gọi mẹ",
    key="Tệp giúp dữ liệu tồn tại lâu dài qua các lần chạy chương trình.")
row(4, L, "Khái niệm Stream và các loại Stream", R, N,
    goal="Hình dung stream là dòng chảy byte tuần tự với một vị trí hiện tại (Position).",
    form="1 cột, 2 tab",
    visual="Ống dẫn nằm ngang nối Tệp và Chương trình. Từng byte (ô nhỏ ghi ký tự và mã số) chảy qua ống. Trên dải byte của tệp có con trỏ Position; nhãn Length. Tab Ghi: ống chảy theo chiều ngược lại.",
    comps="H11, H20", tabs="Đọc (Tệp → Chương trình) | Ghi (Chương trình → Tệp)",
    steps=["Đọc: tệp chứa \"Hi!\" (3 byte: 72, 105, 33). Mở FileStream → Position 0, Length 3.",
           "ReadByte() → 72 ('H'), Position 1. Tiếp → 105 ('i'), → 33 ('!').",
           "ReadByte() lần 4 → −1: đã hết stream.",
           "Thêm một bước: chữ 'Đ' trong UTF-8 chiếm 2 byte (C4 90) để thấy ký tự không luôn bằng 1 byte.",
           "Ghi: WriteByte từng byte, Position và Length cùng tăng."],
    data="\"Hi!\" và 'Đ'",
    key="Stream đọc/ghi lần lượt theo vị trí; hết dữ liệu thì trả −1.",
    cs="FileStream, MemoryStream, NetworkStream đều là Stream: cùng cách đọc/ghi, khác nguồn dữ liệu (nhắc tên ở bước cuối).")
row(4, L, "Tổng quan về các lớp File I/O trong C#", R, K, note="Bài tổng quan. Đề xuất dùng hình sơ đồ tĩnh (nhóm File/Directory/Path, nhóm Stream, nhóm Reader/Writer) thay vì mô phỏng.")
row(4, L, "Ví dụ đơn giản về việc mở và đóng file", R, T,
    goal="Thấy vòng đời của một tệp: mở → dùng → đóng, và trạng thái tệp ở từng bước.",
    form="1 cột, 2 tab",
    visual="Biểu tượng tệp có đèn trạng thái (Đóng xám / Đang mở xanh kèm ổ khóa). Bảng nhỏ 'Tài nguyên hệ điều hành đang giữ'.",
    comps="H08, H12", tabs="Mở rồi đóng | Quên đóng",
    steps=["Mở rồi đóng: new StreamReader(\"a.txt\") → đèn xanh, bảng tài nguyên thêm 1 dòng; ReadLine; Close() → đèn xám, bảng trống.",
           "Quên đóng: không gọi Close → tệp vẫn bị giữ.",
           "Một chương trình khác (ví dụ Notepad) muốn ghi đè tệp → báo tệp đang được dùng.",
           "Bước cuối: tệp chỉ được nhả khi chương trình kết thúc hoặc khi bộ gom rác dọn đối tượng, không biết lúc nào."],
    data="a.txt 2 dòng",
    key="Mỗi lần mở tệp phải có một lần đóng tương ứng.",
    cs="Hiện tượng tệp bị khóa thấy rõ trên Windows; trên macOS/Linux hành vi khóa khác.")
row(4, L, "Thực hành mở và đóng file", A, K, note="Hoạt động đã có nội dung.")

L = "Đọc File Văn bản"
row(4, L, "Giới thiệu lớp StreamReader", R, T,
    goal="Thấy StreamReader đọc từng ký tự và giữ một con trỏ vị trí trong tệp.",
    form="1 cột",
    visual="Trình xem tệp dạng lưới ký tự (mỗi ô một ký tự, ↵ cho xuống dòng). Con trỏ đọc (tam giác) nằm dưới ô hiện tại; ô đã đọc nền xanh nhạt.",
    comps="H12, H11", tabs="1 tab",
    steps=["Tệp chứa \"Hi↵OK\". Mở StreamReader, con trỏ ở ký tự đầu.",
           "Peek() → 72 ('H'), con trỏ không di chuyển.",
           "Read() → 72 ('H'), con trỏ sang 'i'.",
           "Read() → 'i'; Read() → ký tự xuống dòng (10).",
           "Đọc tiếp 'O', 'K'; Read() → −1 và EndOfStream = true.",
           "Xong."],
    data="\"Hi\\nOK\"",
    key="Peek nhìn trước không di chuyển; Read đọc rồi tiến; hết tệp trả −1.",
    cs="Read() trả int (mã ký tự) để có thể trả −1; ép sang char để hiển thị.")
row(4, L, "Đọc từng dòng với ReadLine()", R, N,
    goal="Thấy vòng lặp while ((line = reader.ReadLine()) != null) đọc từng dòng và dừng khi gặp null.",
    form="1 cột",
    visual="Trình xem tệp có số dòng; con trỏ dòng ▶ nhảy từng dòng; ô biến line và count; console in dần. Bước cuối ô line hiện null, con trỏ ở sau dòng cuối.",
    comps="H12", tabs="1 tab",
    steps=["Tệp diem.txt có 4 dòng (dòng 3 trống).",
           "Lần 1: line = \"An,8\", count = 1, in ra.",
           "Lần 2: \"Bình,6\".",
           "Lần 3: line = \"\" (chuỗi rỗng, không phải null) → vòng lặp vẫn tiếp tục.",
           "Lần 4: \"Chi,9\".",
           "Lần 5: line = null → điều kiện sai, thoát vòng lặp. Xong: đã đọc 4 dòng."],
    data="diem.txt: An,8 / Bình,6 / (trống) / Chi,9",
    key="ReadLine trả \"\" cho dòng trống và null khi hết tệp; phải kiểm tra null để dừng.")
row(4, L, "Đọc toàn bộ nội dung với ReadToEnd() và các phương thức đọc khác", R, N,
    goal="So sánh hình dạng kết quả của các cách đọc toàn bộ tệp.",
    form="1 cột, 3 tab",
    visual="Cùng một tệp 3 dòng ở trên. Kết quả ở dưới đổi theo tab: ReadToEnd → một chuỗi dài có ký hiệu ↵; ReadAllLines → mảng 3 ô; ReadLines → băng chuyền đưa từng dòng vào foreach (đọc đến đâu lấy đến đó).",
    comps="H12, H01", tabs="ReadToEnd | File.ReadAllLines | File.ReadLines",
    steps=["ReadToEnd: cả tệp vào một string, độ dài tính cả ký tự xuống dòng.",
           "ReadAllLines: mảng string[3], ký tự xuống dòng bị bỏ.",
           "ReadLines: foreach lấy dòng 1, xử lý xong mới đọc dòng 2... Có thể break sớm mà không đọc phần còn lại.",
           "Bước cuối mỗi tab: khi nào nên dùng."],
    data="Tệp 3 dòng",
    key="Tệp nhỏ: ReadAllText/ReadAllLines cho tiện. Tệp lớn: ReadLine hoặc File.ReadLines để không nạp hết vào bộ nhớ.",
    cs="Xuống dòng là \\r\\n trên Windows, \\n trên macOS/Linux.")
row(4, L, "Đọc nội dung file và hiển thị", R, K, note="Trùng ý 'Đọc từng dòng với ReadLine()'. Đính kèm lại mô phỏng đó.")
row(4, L, "Đọc file và xử lý dữ liệu đơn giản", R, T,
    goal="Kết hợp đọc từng dòng, chuyển đổi an toàn bằng TryParse và tính toán.",
    form="1 cột",
    visual="Trình xem tệp numbers.txt; mỗi dòng đi qua 'bộ chuyển đổi' TryParse: ✓ cộng vào tổng, ✗ thả vào danh sách dòng lỗi. Ô Tổng, Max, Số dòng lỗi cập nhật mỗi bước.",
    comps="H12, H21", tabs="1 tab",
    steps=["Dòng 1 \"12\" → tổng 12, max 12.",
           "Dòng 2 \"7\" → tổng 19.",
           "Dòng 3 \"abc\" → TryParse false, ghi lỗi dòng 3.",
           "Dòng 4 \"20\" → tổng 39, max 20.",
           "Dòng 5 \"\" → lỗi dòng 5.",
           "Dòng 6 \"5\" → tổng 44. Xong: tổng 44, max 20, 2 dòng lỗi (3 và 5)."],
    data="numbers.txt: 12, 7, abc, 20, (trống), 5",
    key="Dữ liệu trong tệp luôn là chuỗi; phải chuyển đổi an toàn và ghi lại dòng lỗi.")
row(4, L, "Thực hành: Đọc file log và lọc thông tin", A, K, note="Hoạt động đã có nội dung.")

L = "Ghi File Văn bản"
row(4, L, "Giới thiệu lớp StreamWriter", R, N,
    goal="Thấy StreamWriter ghi vào bộ đệm trước, và chỉ đẩy xuống tệp khi Flush/Close/Dispose.",
    form="2 cột song song: Quên đóng | Có using",
    visual="Ba khối Chương trình → Bộ đệm (hộp có vạch đầy dần) → Tệp trên đĩa (hiện kích thước byte). Khi Close/Dispose, nội dung bộ đệm 'đổ' xuống tệp.",
    comps="H13, H12", tabs="1 tab",
    steps=["Hai cột: new StreamWriter(\"log.txt\").",
           "WriteLine(\"Dòng 1\"), WriteLine(\"Dòng 2\") → chữ nằm trong bộ đệm, tệp vẫn 0 byte.",
           "Trái: chương trình kết thúc mà không Close → bộ đệm bị bỏ, tệp vẫn 0 byte.",
           "Phải: rời khối using → Dispose → Flush → tệp có 2 dòng.",
           "Xong: so sánh nội dung tệp hai bên."],
    data="2 dòng log",
    key="Luôn đóng StreamWriter, tốt nhất bằng using, để dữ liệu thật sự được ghi xuống tệp.",
    cs="Bộ đệm có kích thước giới hạn; ghi nhiều dữ liệu thì một phần tự được đẩy xuống. Mô phỏng dùng dữ liệu nhỏ để thấy rõ hiện tượng mất dữ liệu.")
row(4, L, "Ghi từng dòng và thêm mới nội dung", R, N,
    goal="Phân biệt Write và WriteLine, và phân biệt chế độ ghi đè với ghi nối.",
    form="1 cột, 3 tab",
    visual="Trình xem tệp 'trước' và 'sau' đặt cạnh nhau; dòng cũ bị gạch ngang khi ghi đè; dòng mới tô xanh; con trỏ ghi | cho thấy Write không xuống dòng.",
    comps="H12", tabs="Write và WriteLine | Ghi đè (append: false) | Ghi nối (append: true)",
    steps=["Write/WriteLine: Write(\"A\"); Write(\"B\"); WriteLine(\"C\"); WriteLine(\"D\") → tệp \"ABC↵D↵\"; con trỏ ghi minh họa từng bước.",
           "Ghi đè: tệp cũ có 2 dòng; new StreamWriter(path) → tệp bị làm rỗng ngay khi mở; ghi 1 dòng → còn 1 dòng.",
           "Ghi nối: new StreamWriter(path, append: true) → con trỏ đặt ở cuối; ghi thêm 1 dòng → 3 dòng.",
           "Bước cuối mỗi tab: so sánh trước/sau."],
    data="Tệp cũ: Dòng cũ 1, Dòng cũ 2",
    key="Mặc định StreamWriter ghi đè; muốn giữ nội dung cũ dùng append: true hoặc File.AppendAllText.")
row(4, L, "Thực hành ghi dữ liệu vào file mới", R, K, note="Trùng ý 'Ghi từng dòng và thêm mới nội dung'. Đính kèm lại.")
row(4, L, "Thêm dữ liệu vào cuối file", A, K, note="Hoạt động đã có nội dung.")
row(4, L, "Thực hành tạo file cấu hình đơn giản", R, T,
    goal="Lưu cấu hình dạng key=value ra tệp và đọc lại thành Dictionary.",
    form="1 cột, 2 tab",
    visual="Tủ ngăn Dictionary bên trái ↔ tệp config.txt bên phải; mũi tên nối từng cặp với từng dòng. Dòng ghi chú và dòng sai định dạng tô xám 'bỏ qua'.",
    comps="H02, H12", tabs="Ghi cấu hình | Đọc cấu hình",
    steps=["Ghi: foreach (var kv in settings) WriteLine($\"{kv.Key}={kv.Value}\") → 3 dòng.",
           "Đọc: ReadAllLines; dòng \"# Cấu hình\" bắt đầu bằng # → bỏ qua.",
           "Dòng \"theme=dark\" → Split('=', 2) → [\"theme\", \"dark\"] → thêm vào Dictionary.",
           "Dòng \"url=https://a.vn/?x=1\" → Split('=', 2) giữ nguyên phần sau dấu = đầu tiên.",
           "Dòng \"abc\" không có '=' → bỏ qua. Xong."],
    data="theme=dark, fontSize=14, url=https://a.vn/?x=1",
    key="Ghi và đọc phải cùng một quy ước định dạng.",
    cs="Split('=', 2) chỉ tách ở dấu '=' đầu tiên để giá trị được phép chứa '='.")
row(4, L, "Ghi kết quả tính toán vào file", A, K, note="Hoạt động đã có nội dung.")

L = "Quản lý Đường dẫn và Thư mục"
row(4, L, "Lớp Path: Thao tác với đường dẫn file", R, N,
    goal="Tách và ghép đường dẫn an toàn bằng Path thay vì cộng chuỗi.",
    form="1 cột, 3 tab",
    visual="Đường dẫn hiển thị như dãy chip: ổ đĩa | thư mục | thư mục | tên tệp | phần mở rộng. Mỗi phương thức tô các chip mà nó trả về.",
    comps="H14", tabs="Tách đường dẫn | Ghép đường dẫn | Đổi phần mở rộng",
    steps=["Tách: C:\\Data\\Reports\\bao-cao.txt → GetDirectoryName = C:\\Data\\Reports; GetFileName = bao-cao.txt; GetFileNameWithoutExtension = bao-cao; GetExtension = .txt.",
           "Ghép: \"Data\" + \"Reports\" → \"DataReports\" (sai, thiếu dấu phân cách); Path.Combine(\"Data\", \"Reports\", \"a.txt\") → Data\\Reports\\a.txt.",
           "Ghép với đường dẫn tuyệt đối: Path.Combine(\"Data\", @\"C:\\Temp\") → C:\\Temp (phần trước bị bỏ).",
           "Đổi đuôi: Path.ChangeExtension(\"bao-cao.txt\", \".csv\") → bao-cao.csv (chỉ đổi chuỗi, không đổi tệp thật).",
           "Bước cuối mỗi tab: kết quả."],
    data="C:\\Data\\Reports\\bao-cao.txt",
    key="Path xử lý dấu phân cách và các trường hợp đặc biệt thay cho việc cộng chuỗi.",
    cs="Ví dụ dùng đường dẫn kiểu Windows; trên macOS/Linux dấu phân cách là '/', Path.Combine tự dùng Path.DirectorySeparatorChar.")
row(4, L, "Lớp Directory: Quản lý thư mục", R, N,
    goal="Thấy các thao tác thư mục tác động lên cây thư mục ra sao.",
    form="1 cột, 4 tab",
    visual="Cây thư mục có biểu tượng thư mục/tệp và thụt lề; nút mới tô xanh, nút bị xóa mờ đỏ, nút được liệt kê viền vàng; vị trí lỗi gắn dấu chấm than.",
    comps="H15, H03", tabs="CreateDirectory | GetFiles | Delete | Exists",
    steps=["CreateDirectory(@\"Data\\2024\\Thang10\") → tạo cả các cấp còn thiếu; gọi lại khi đã có → không lỗi.",
           "GetFiles(\"Data\", \"*.txt\") → chỉ tệp .txt ở cấp Data; thêm SearchOption.AllDirectories → cả thư mục con.",
           "Delete(\"Data\") khi bên trong còn nội dung → IOException; Delete(\"Data\", true) → xóa cả cây.",
           "Exists(\"Data\") trước và sau khi xóa: true → false.",
           "Bước cuối mỗi tab: trạng thái cây."],
    data="Cây: Data/{a.txt, b.csv, 2024/{c.txt}}",
    key="Directory làm việc với thư mục; xóa thư mục không rỗng cần recursive: true.",
    cs="Thông điệp lỗi cụ thể khác nhau theo hệ điều hành; mô phỏng ghi tên loại exception.")
row(4, L, "Các phương thức hữu ích của lớp File", R, T,
    goal="Nắm nhanh Exists, Copy, Move, Delete và cách chúng phản ứng khi tệp đích đã có hoặc tệp không tồn tại.",
    form="1 cột, 4 tab",
    visual="Cây thư mục nhỏ; thao tác Copy/Move có mũi tên cong từ nguồn tới đích; lỗi gắn thẻ đỏ.",
    comps="H15, H03", tabs="Exists | Copy | Move | Delete",
    steps=["Exists(\"a.txt\") → true; Exists(\"z.txt\") → false.",
           "Copy(\"a.txt\", \"b.txt\") khi b.txt đã có → IOException; Copy(..., overwrite: true) → ghi đè.",
           "Move(\"a.txt\", @\"Backup\\a.txt\") → tệp chuyển chỗ (cũng dùng để đổi tên).",
           "Delete(\"z.txt\") khi tệp không tồn tại → không lỗi; Delete(@\"KhongCo\\z.txt\") khi thư mục không có → DirectoryNotFoundException.",
           "Bước cuối mỗi tab: trạng thái cây."],
    data="Cây: a.txt, b.txt, Backup/",
    key="Biết trước phương thức nào ném lỗi giúp chọn chỗ đặt try-catch hoặc kiểm tra trước.")
row(4, L, "Thực hành: Tạo và xóa thư mục", A, K, note="Hoạt động đã có nội dung.")
row(4, L, "Kiểm tra sự tồn tại của file/thư mục", A, K, note="Hoạt động đã có nội dung.")
row(4, L, "Bài tập: Quản lý Thư mục và File Tự động", A, K, note="Bài tập nộp, giữ nguyên.")

L = "Xử lý Lỗi và Tài nguyên trong File I/O"
row(4, L, "Các loại ngoại lệ phổ biến trong File I/O", R, N,
    goal="Nhận ra exception nào xảy ra ở tình huống tệp/thư mục nào, và quan hệ giữa chúng.",
    form="1 cột, 4 tab",
    visual="Cây thư mục đánh dấu vị trí gây lỗi (nút đỏ; ổ khóa cho tệp không có quyền hoặc đang bị dùng). Thẻ exception kèm cách xử lý. Bước cuối: cây phân cấp nhỏ IOException → FileNotFoundException, DirectoryNotFoundException.",
    comps="H15, H03, H06", tabs="FileNotFoundException | DirectoryNotFoundException | UnauthorizedAccessException | IOException (tệp đang được dùng)",
    steps=["FileNotFound: đọc data.txt không tồn tại.",
           "DirectoryNotFound: ghi vào KhongCo\\a.txt khi thư mục KhongCo không tồn tại.",
           "UnauthorizedAccess: ghi vào tệp chỉ đọc hoặc thư mục hệ thống.",
           "IOException: tệp đang được chương trình khác mở để ghi.",
           "Bước cuối: catch (IOException) bắt được cả FileNotFound và DirectoryNotFound; UnauthorizedAccessException không phải IOException nên cần catch riêng."],
    data="Tình huống cố định",
    key="Bắt cụ thể trước, IOException sau; nhớ UnauthorizedAccessException là nhánh riêng.",
    cs="FileNotFoundException, DirectoryNotFoundException kế thừa IOException; UnauthorizedAccessException kế thừa SystemException. Tình huống tệp bị khóa thấy rõ trên Windows.")
row(4, L, "Sử dụng khối try-catch-finally", R, K, note="Trùng ý 'Sử dụng finally để đảm bảo tài nguyên được giải phóng' (chương 2). Đính kèm lại.")
row(4, L, "Câu lệnh 'using' để quản lý tài nguyên (IDisposable)", R, K, note="Trùng ý 'using statement như một cách thay thế cho finally' (chương 3). Đính kèm lại.")
row(4, L, "Xử lý lỗi khi file không tồn tại", R, T,
    goal="So sánh hai cách: kiểm tra File.Exists trước và bắt FileNotFoundException.",
    form="2 cột song song: Kiểm tra trước | Bắt exception",
    visual="Cây thư mục chung ở trên; đường đi code của mỗi cột tô theo tình huống.",
    comps="H15, H04", tabs="Tệp có tồn tại | Tệp không tồn tại",
    steps=["Tệp có: cả hai đọc được nội dung.",
           "Tệp không có: trái File.Exists trả false → báo \"Không tìm thấy tệp\", không ném lỗi.",
           "Phải: ReadAllText ném FileNotFoundException → catch báo cùng thông điệp.",
           "Bước cuối: kiểm tra trước giúp code rõ ràng, nhưng vẫn nên có try-catch vì tệp có thể bị xóa ngay sau khi kiểm tra."],
    data="data.txt",
    key="Kết hợp cả hai: kiểm tra trước cho trường hợp thường gặp, try-catch cho trường hợp bất ngờ.")
row(4, L, "Thực hành: Viết chương trình đọc/ghi file có xử lý lỗi đầy đủ", A, K, note="Hoạt động đã có nội dung.")
row(4, L, "Viết code an toàn với 'using'", A, K, note="Hoạt động đã có nội dung.")

# ======================= CHƯƠNG 5: CSV =======================
L = "Giới thiệu về CSV và Cấu trúc cơ bản"
row(5, L, "CSV là gì và tại sao lại quan trọng?", R, N,
    goal="Thấy bảng tính và tệp CSV là hai cách nhìn của cùng một dữ liệu.",
    form="1 cột (hai khung hình minh họa cạnh nhau)",
    visual="Trái: bảng kiểu Excel (lưới ô, hàng tiêu đề đậm). Phải: văn bản CSV thô. Mỗi bước tô cùng màu một hàng ở hai bên; dấu phẩy được tô như đường kẻ ô.",
    comps="H16", tabs="1 tab",
    steps=["Hàng tiêu đề Ma | Ten | Gia ↔ dòng 1 \"Ma,Ten,Gia\".",
           "Bản ghi 1 ↔ \"SP01,Bút bi,5000\" (dấu phẩy = ranh giới cột).",
           "Bản ghi 2 ↔ \"SP02,Vở,12000\".",
           "Bản ghi 3 ↔ \"SP03,Thước,8000\".",
           "Xong: mở bằng Notepad thấy văn bản, mở bằng Excel thấy bảng."],
    data="Ma,Ten,Gia / SP01,Bút bi,5000 / SP02,Vở,12000 / SP03,Thước,8000",
    key="CSV là văn bản thuần: mỗi dòng một bản ghi, dấu phẩy ngăn cách trường; người, Excel và chương trình đều đọc được.")
row(5, L, "Cấu trúc cơ bản của File CSV", R, T,
    goal="Gọi đúng tên từng thành phần của tệp CSV.",
    form="1 cột",
    visual="Văn bản CSV với các nhãn chú thích bật lên chỉ vào từng phần: Hàng tiêu đề, Bản ghi, Trường, Dấu phân cách, Ký tự xuống dòng (↵), Trường rỗng (,,).",
    comps="H16", tabs="1 tab",
    steps=["Nhãn 'Hàng tiêu đề'.", "Nhãn 'Bản ghi' (một dòng).", "Nhãn 'Trường' (một giá trị).",
           "Nhãn 'Dấu phân cách' (dấu phẩy).", "Nhãn 'Trường rỗng': SP04,,3000 (thiếu tên).",
           "Nhãn 'Xuống dòng' ↵. Xong: ghi chú một số vùng (như Việt Nam) Excel xuất CSV dùng dấu ';'."],
    data="Như bài trước, thêm SP04,,3000",
    key="Hiểu đúng tên gọi giúp đọc tài liệu và viết code xử lý CSV chính xác.",
    cs="Excel ở vùng dùng dấu phẩy làm dấu thập phân thường xuất CSV với dấu ';'.")
row(5, L, "Ví dụ minh họa cấu trúc CSV", R, K, note="Trùng ý 'CSV là gì'. Đính kèm lại mô phỏng đó.")
row(5, L, "Thực hành: Phân tích cấu trúc file CSV", A, K, note="Hoạt động đã có nội dung.")

L = "Đọc dữ liệu từ File CSV"
row(5, L, "Các bước cơ bản để đọc File CSV trong C#", R, N,
    goal="Thấy các bước biến một dòng văn bản thành một đối tượng trong List.",
    form="1 cột",
    visual="Dây chuyền ngang: Dòng văn bản → Split(',') → mảng 3 ô → chuyển kiểu (int.Parse giá) → thẻ đối tượng Product → thả vào List<Product>. Mỗi dòng lần lượt đi qua các trạm.",
    comps="H17, H12, H01", tabs="1 tab",
    steps=["File.ReadAllLines(\"sp.csv\") → mảng 4 dòng.",
           "Bỏ dòng tiêu đề (bắt đầu từ i = 1).",
           "Dòng 1: Split(',') → [\"SP01\", \"Bút bi\", \"5000\"].",
           "int.Parse(\"5000\") → 5000; tạo new Product(\"SP01\", \"Bút bi\", 5000); Add vào List.",
           "Lặp lại cho dòng 2 và 3.",
           "Xong: List có 3 Product."],
    data="sp.csv như bài 'CSV là gì'",
    key="Đọc CSV = đọc dòng → tách trường → chuyển kiểu → tạo đối tượng.")
row(5, L, "Đọc file CSV đơn giản (không có dấu nháy kép)", R, K, note="Trùng ý 'Các bước cơ bản để đọc File CSV'. Đính kèm lại.")
row(5, L, "Thực hành: Đọc CSV không dấu nháy kép", A, K, note="Hoạt động đã có nội dung.")
row(5, L, "Xử lý các trường có dấu nháy kép và dấu phân cách đặc biệt", R, N,
    goal="Thấy vì sao Split(',') sai với trường chứa dấu phẩy, và bộ phân tích theo trạng thái xử lý đúng ra sao.",
    form="2 cột song song: Split(',') | Bộ phân tích theo trạng thái",
    visual="Trái: dòng bị cắt ở mọi dấu phẩy → 4 ô, ô thừa tô đỏ. Phải: con trỏ chạy trên từng ký tự; huy hiệu 'Trong nháy: Có/Không'; dấu phẩy trong nháy tô xám 'không cắt'; trường hoàn tất được thả vào ô kết quả.",
    comps="H18, H17", tabs="Dấu phẩy trong trường | Nháy kép trong trường",
    steps=["Tab 1, dòng SP04,\"Bút, mực xanh\",15000.",
           "Trái: Split → [\"SP04\", \"\\\"Bút\", \" mực xanh\\\"\", \"15000\"] → 4 trường, sai.",
           "Phải (mỗi bước là một sự kiện): gặp dấu phẩy ngoài nháy → chốt \"SP04\"; gặp nháy mở → Trong nháy: Có; gặp dấu phẩy trong nháy → giữ lại; gặp nháy đóng → Trong nháy: Không; gặp dấu phẩy → chốt \"Bút, mực xanh\"; hết dòng → chốt \"15000\".",
           "Tab 2, dòng SP05,\"Sách \"\"C#\"\" cơ bản\",90000: hai nháy liền nhau trong trường được hiểu là một ký tự nháy → Sách \"C#\" cơ bản.",
           "Xong: 3 trường đúng."],
    data="SP04,\"Bút, mực xanh\",15000 và SP05,\"Sách \"\"C#\"\" cơ bản\",90000",
    key="Trường chứa dấu phẩy, nháy kép hoặc xuống dòng được bọc trong nháy kép; nháy bên trong viết thành \"\" (theo RFC 4180).",
    cs="Mô phỏng nhóm các ký tự thường thành một bước để không quá dài; chỉ dừng ở ký tự đặc biệt (nháy, dấu phẩy).")
row(5, L, "Phân tích CSV có dấu nháy kép", A, K, note="Hoạt động đã có nội dung.")
row(5, L, "Thực hành: Đọc CSV với dấu nháy kép và chuyển đổi dữ liệu", A, K, note="Hoạt động đã có nội dung.")

L = "Ghi dữ liệu ra File CSV"
row(5, L, "Các bước cơ bản để ghi File CSV trong C#", R, N,
    goal="Thấy chiều ngược lại: đối tượng → dòng văn bản → tệp.",
    form="1 cột",
    visual="Dây chuyền ngược: thẻ Product → mảng giá trị → string.Join(\",\") → dòng → trình xem tệp đầy dần. Dòng tiêu đề được ghi đầu tiên.",
    comps="H17, H12", tabs="1 tab",
    steps=["Mở StreamWriter trong using.",
           "WriteLine(\"Ma,Ten,Gia\") → dòng tiêu đề.",
           "Product 1 → [\"SP01\", \"Bút bi\", \"5000\"] → Join → \"SP01,Bút bi,5000\" → ghi.",
           "Lặp với 2 sản phẩm còn lại.",
           "Rời using → tệp được đóng và có 4 dòng. Xong."],
    data="3 Product",
    key="Ghi CSV = ghi tiêu đề → với mỗi đối tượng: lấy giá trị → nối bằng dấu phẩy → ghi dòng.",
    cs="Số thập phân nên định dạng bằng CultureInfo.InvariantCulture để 12.5 không thành \"12,5\" (làm lệch cột) trên máy cài tiếng Việt.")
row(5, L, "Ghi file CSV đơn giản", R, K, note="Trùng ý 'Các bước cơ bản để ghi File CSV'. Đính kèm lại.")
row(5, L, "Thực hành: Ghi CSV không dấu nháy kép", A, K, note="Hoạt động đã có nội dung.")
row(5, L, "Quy tắc xử lý dấu nháy kép khi ghi CSV", R, N,
    goal="Áp dụng đúng quy tắc 'bọc nháy và nhân đôi nháy' khi ghi một trường.",
    form="1 cột, 4 tab theo giá trị trường",
    visual="Thẻ 'giá trị trường' đi qua 2 trạm: 'Có dấu phẩy / nháy / xuống dòng?' và 'Nhân đôi dấu nháy'. Trạm sáng/tắt; kết quả cuối hiện trong ô CSV có tô màu phần được thêm.",
    comps="H09, H17", tabs="Bút bi | Hà Nội, Việt Nam | Sách \"C#\" | Có xuống dòng",
    steps=["Bút bi → không có ký tự đặc biệt → giữ nguyên: Bút bi.",
           "Hà Nội, Việt Nam → có dấu phẩy → bọc nháy: \"Hà Nội, Việt Nam\".",
           "Sách \"C#\" → có nháy → nhân đôi nháy rồi bọc: \"Sách \"\"C#\"\"\".",
           "Dòng 1↵Dòng 2 → có xuống dòng → bọc nháy.",
           "Bước cuối mỗi tab: hiện hàm EscapeCsv và kết quả."],
    data="4 giá trị mẫu",
    key="Ghi đúng quy tắc thì đọc lại bằng bộ phân tích chuẩn sẽ ra đúng giá trị ban đầu.",
    cs="EscapeCsv: nếu chứa , \" \\n hoặc \\r thì trả \"\\\"\" + s.Replace(\"\\\"\", \"\\\"\\\"\") + \"\\\"\".")
row(5, L, "Thực hành: Xây dựng dòng CSV từ dữ liệu", A, K, note="Hoạt động đã có nội dung.")
row(5, L, "Bài tập: Xuất dữ liệu sản phẩm ra định dạng CSV", A, K, note="Bài tập nộp, giữ nguyên.")

L = "Xử lý CSV nâng cao và Ứng dụng thực tế"
row(5, L, "Xử lý các trường hợp đặc biệt và lỗi trong CSV", R, N,
    goal="Đọc một tệp CSV 'bẩn' mà không sập: giữ bản ghi hợp lệ và báo lỗi kèm số dòng.",
    form="1 cột",
    visual="Trình xem tệp; mỗi dòng đi qua các kiểm tra (dòng trống? đủ 3 cột? giá là số?). Hai thùng bên dưới: Hợp lệ (thẻ xanh) và Lỗi (thẻ đỏ ghi số dòng và lý do).",
    comps="H12, H09, H21", tabs="1 tab",
    steps=["Dòng 1: tiêu đề, bỏ qua.",
           "Dòng 2 SP01,Bút,5000 → hợp lệ.",
           "Dòng 3 trống → bỏ qua (không tính lỗi).",
           "Dòng 4 SP02,Vở → thiếu cột → thùng Lỗi: \"Dòng 4: cần 3 cột, có 2\".",
           "Dòng 5 SP03,Thước,abc → giá không phải số → thùng Lỗi: \"Dòng 5: giá không hợp lệ\".",
           "Dòng 6 SP04,Tẩy,3000 → hợp lệ. Xong: 2 hợp lệ, 2 lỗi."],
    data="6 dòng như kịch bản",
    key="Không để một dòng hỏng làm hỏng cả tệp; ghi lại số dòng để người dùng sửa.")
row(5, L, "Xây dựng một lớp CsvHelper đơn giản", R, T,
    goal="Thấy lợi ích gom code đọc/ghi CSV vào một lớp: ghi rồi đọc lại phải ra đúng dữ liệu (round-trip).",
    form="1 cột",
    visual="Hộp lớp CsvHelper (ToCsvLine, ParseLine) ở trên; dưới là vòng tròn: đối tượng → dòng CSV → đối tượng mới; bảng so sánh từng trường với dấu ✓.",
    comps="H17, H19", tabs="1 tab",
    steps=["Product gốc: SP09, Bút, mực \"đỏ\", 15000.",
           "CsvHelper.ToCsvLine → SP09,\"Bút, mực \"\"đỏ\"\"\",15000.",
           "CsvHelper.ParseLine → Product mới.",
           "So sánh từng trường: Ma ✓, Ten ✓, Gia ✓.",
           "Xong."],
    data="Product SP09",
    key="Một lớp, hai chiều, kiểm tra bằng round-trip: ghi rồi đọc lại phải giống ban đầu.")
row(5, L, "Viết hàm đọc CSV linh hoạt", R, T,
    goal="Đọc đúng dữ liệu kể cả khi thứ tự cột trong tệp thay đổi, nhờ ánh xạ tên cột → chỉ số.",
    form="2 cột song song: Đọc theo chỉ số cố định | Đọc theo tên cột",
    visual="Hàng tiêu đề với các thẻ cột; bên phải dựng Dictionary 'Ten → 0, Gia → 1, Ma → 2' từ tiêu đề, mũi tên nối từ tên tới chỉ số.",
    comps="H02, H17, H03", tabs="1 tab",
    steps=["Tệp có thứ tự cột mới: Ten,Gia,Ma.",
           "Trái giả định parts[0] là Mã → Mã = \"Bút bi\" (sai).",
           "Trái đọc giá ở parts[2] → int.Parse(\"SP01\") → FormatException.",
           "Phải đọc tiêu đề → dựng Dictionary tên cột → chỉ số.",
           "Phải lấy giá trị theo col[\"Ma\"], col[\"Gia\"] → đúng.",
           "Xong."],
    data="Ten,Gia,Ma / Bút bi,5000,SP01",
    key="Đọc theo tên cột giúp chương trình chịu được thay đổi thứ tự cột.")
row(5, L, "Bài tập: Xây dựng Công cụ Quản lý Dữ liệu Sản phẩm CSV", A, K, note="Bài tập nộp, giữ nguyên.")
row(5, L, "Thảo luận về ưu nhược điểm của CSV so với các định dạng khác", A, K, note="Thảo luận, giữ nguyên.")

# ======================= CHƯƠNG 6: SERIALIZATION =======================
L = "Giới Thiệu về Serialization và Khái Niệm Cơ Bản"
row(6, L, "Khái niệm Serialization là gì?", R, N,
    goal="Hiểu serialize là biến đối tượng thành dạng lưu/gửi được, deserialize là dựng lại đối tượng từ dạng đó.",
    form="1 cột, 2 tab",
    visual="Thẻ đối tượng Person#1 trong khung RAM → 'máy ép' Serialize → chuỗi JSON → biểu tượng tệp trên ổ đĩa. Chiều ngược lại: 'máy dựng' Deserialize → Person#2.",
    comps="H19, H10", tabs="Serialize | Deserialize",
    steps=["Serialize: Person { Name = \"An\", Age = 20 } (Person#1).",
           "JsonSerializer.Serialize(p) → {\"Name\":\"An\",\"Age\":20}.",
           "File.WriteAllText(\"p.json\", json) → tệp xuất hiện trên ổ đĩa.",
           "Deserialize: tắt chương trình, RAM trống.",
           "ReadAllText → Deserialize<Person> → Person#2: đối tượng mới, cùng dữ liệu.",
           "Xong."],
    data="Person An, 20",
    key="Dữ liệu được bảo toàn qua serialize/deserialize, nhưng đối tượng nhận lại là đối tượng mới.")
row(6, L, "Tại sao cần Serialization?", R, K, note="Bài lý thuyết; ý chính đã có trong 'Khái niệm Serialization'.")
row(6, L, "Các loại Serialization phổ biến trong C#", R, N,
    goal="Nhìn cùng một đối tượng ở 3 định dạng và so sánh độ dễ đọc, kích thước.",
    form="1 cột, 3 tab",
    visual="Thẻ đối tượng cố định ở trên; khung kết quả đổi theo tab: JSON (tô màu khóa/giá trị), XML (cây thẻ mở/đóng), Nhị phân (dải ô hex, không đọc được). Thanh kích thước so sánh số byte, tính từ dữ liệu thật.",
    comps="H19, H20", tabs="JSON | XML | Nhị phân",
    steps=["JSON: {\"Id\":7,\"Name\":\"Bút\",\"Price\":5000}, ghi số byte.",
           "XML: <Product><Id>7</Id><Name>Bút</Name><Price>5000</Price></Product>, ghi số byte.",
           "Nhị phân: 13 byte do BinaryWriter ghi (xem bài Binary Serialization).",
           "Bước cuối: bảng so sánh dễ đọc / kích thước / dùng khi nào."],
    data="Product { Id = 7, Name = \"Bút\", Price = 5000 }",
    key="JSON gọn và dễ đọc (phổ biến nhất); XML dài hơn nhưng chặt chẽ; nhị phân gọn nhất nhưng người không đọc được.",
    cs="CẢNH BÁO: BinaryFormatter đã lỗi thời từ .NET 5 và bị loại bỏ trong .NET 9 (gọi sẽ ném exception) vì rủi ro bảo mật. Tab nhị phân dùng BinaryWriter để minh họa.",
    note="Học liệu chương này dựa nhiều vào BinaryFormatter: đề xuất cập nhật nội dung bài.")
row(6, L, "Cấu trúc lớp cần thiết cho Serialization", A, K, note="Hoạt động đã có nội dung.")

L = "Binary Serialization"
row(6, L, "Giới thiệu Binary Serialization", R, T,
    goal="Thấy dữ liệu nhị phân là các byte được ghi theo đúng thứ tự trường, và phải đọc lại đúng thứ tự đó.",
    form="1 cột, 2 tab",
    visual="Dải byte (ô hex) đầy dần khi ghi từng trường, mỗi trường một màu: Id 4 byte, Name gồm 1 byte độ dài + các byte UTF-8, Price 4 byte. Tab Đọc có con trỏ chạy trên dải byte.",
    comps="H20, H11", tabs="Ghi (BinaryWriter) | Đọc (BinaryReader)",
    steps=["Ghi Id = 7 → 07 00 00 00 (int 4 byte, byte thấp trước).",
           "Ghi Name = \"Bút\" → 04 42 C3 BA 74 (độ dài 4 byte, 'ú' chiếm 2 byte UTF-8).",
           "Ghi Price = 5000 → 88 13 00 00. Tổng 13 byte.",
           "Đọc đúng thứ tự ReadInt32, ReadString, ReadInt32 → khôi phục đúng.",
           "Đọc sai thứ tự (ReadString trước) → dữ liệu vô nghĩa.",
           "Xong."],
    data="Product { Id = 7, Name = \"Bút\", Price = 5000 }",
    key="Định dạng nhị phân gọn nhưng không tự mô tả: người đọc phải biết đúng thứ tự và kiểu của từng trường.",
    cs="Byte được tính đúng theo BinaryWriter (little-endian, độ dài chuỗi mã hóa 7-bit). Không dùng BinaryFormatter (đã bị loại bỏ trong .NET 9).",
    note="Đề xuất cập nhật học liệu: thay BinaryFormatter bằng BinaryWriter/BinaryReader hoặc JSON.")
row(6, L, "Attribut `[Serializable]` và `BinaryFormatter`", R, K,
    note="Không làm mô phỏng: BinaryFormatter đã lỗi thời và bị loại bỏ trong .NET 9. Đề xuất cập nhật nội dung bài (cảnh báo bảo mật, hướng dẫn chuyển sang System.Text.Json).")
row(6, L, "Thực hành: Viết chương trình Serialization đối tượng đơn giản", A, K, note="Hoạt động đã có nội dung (cần rà soát nếu dùng BinaryFormatter).")
row(6, L, "Thực hành Deserialize đối tượng từ file", A, K, note="Hoạt động đã có nội dung (cần rà soát nếu dùng BinaryFormatter).")
row(6, L, "Bài tập: Xây dựng hệ thống lưu trữ thông tin sản phẩm", A, K, note="Bài tập nộp, giữ nguyên.")

L = "XML Serialization"
row(6, L, "Giới thiệu XML Serialization", R, N,
    goal="Thấy XmlSerializer ánh xạ thuộc tính public thành phần tử XML, và điều kiện lớp cần có.",
    form="1 cột, 3 tab",
    visual="Thẻ đối tượng bên trái, cây XML bên phải; đường nối màu từ mỗi thuộc tính public tới thẻ XML tương ứng; trường private không có đường nối (mờ).",
    comps="H19, H03", tabs="Serialize | Deserialize | Thiếu constructor không tham số",
    steps=["Serialize: new XmlSerializer(typeof(Product)); mỗi thuộc tính public → một phần tử (<Id>, <Name>, <Price>); trường private bị bỏ qua.",
           "Deserialize: đọc XML → tạo Product mới bằng constructor không tham số → gán từng thuộc tính.",
           "Thiếu constructor không tham số: new XmlSerializer(typeof(Product)) → InvalidOperationException \"Product cannot be serialized because it does not have a parameterless constructor.\"",
           "Bước cuối mỗi tab: kết quả."],
    data="Product { Id = 7, Name = \"Bút\", Price = 5000 }",
    key="XmlSerializer cần lớp public, constructor không tham số, và chỉ ghi các thành viên public.")
row(6, L, "Bài thực hành: Sử dụng XmlSerializer để Serialize và Deserialize", A, K, note="Hoạt động đã có nội dung.")
row(6, L, "Quy trình Thực thi Serialization và Deserialization với XmlSerializer", R, K, note="Trùng ý 'Giới thiệu XML Serialization'. Đính kèm lại.")
row(6, L, "Xử lý các thuộc tính và phần tử trong XML", A, K, note="Hoạt động đã có nội dung.")
row(6, L, "Bài tập: Xây dựng Hệ thống Lưu trữ Cấu hình Người dùng bằng XML", A, K, note="Bài tập nộp, giữ nguyên.")

L = "JSON Serialization với System.Text.Json"
row(6, L, "Giới thiệu JSON Serialization trong C#", R, N,
    goal="Đọc được JSON và thấy kiểu dữ liệu C# tương ứng với kiểu JSON.",
    form="1 cột, 3 tab",
    visual="Thẻ đối tượng ↔ JSON tô màu theo kiểu (chuỗi xanh lá, số cam, bool tím, null xám); ngoặc [ ] và { } được làm nổi; bảng chú giải kiểu bên cạnh.",
    comps="H19", tabs="Đối tượng đơn | Có danh sách | Đối tượng lồng nhau",
    steps=["Đối tượng đơn: Name (string) → \"An\", Age (int) → 20, IsActive (bool) → true, Email (null) → null.",
           "Có danh sách: List<string> Skills → [\"C#\", \"SQL\"].",
           "Lồng nhau: Address { City = \"Hà Nội\" } → \"Address\": { \"City\": \"Hà Nội\" }.",
           "Bước cuối mỗi tab: JSON hoàn chỉnh."],
    data="User mẫu",
    key="string → \"...\", số → số, bool → true/false, List → [ ], đối tượng con → { }.",
    cs="Mặc định System.Text.Json mã hóa ký tự có dấu thành \\uXXXX; muốn giữ tiếng Việt dễ đọc dùng JavaScriptEncoder.UnsafeRelaxedJsonEscaping (ghi chú ở bước cuối).")
row(6, L, "Sử dụng `System.Text.Json` cơ bản", R, N,
    goal="Biết 3 tùy chọn hay dùng và cái bẫy phân biệt chữ hoa thường khi đọc JSON.",
    form="1 cột, 3 tab",
    visual="Hai khung JSON trước/sau khi đổi tùy chọn đặt cạnh nhau. Tab 3: thẻ đối tượng sau Deserialize có trường bị null/0 tô đỏ, sau khi bật tùy chọn chuyển xanh.",
    comps="H19", tabs="WriteIndented | Tên camelCase | Đọc JSON chữ thường",
    steps=["WriteIndented = false → một dòng; = true → xuống dòng, thụt lề.",
           "PropertyNamingPolicy = JsonNamingPolicy.CamelCase → \"Name\" thành \"name\".",
           "Đọc {\"name\":\"An\",\"age\":20} vào Person (Name, Age) với tùy chọn mặc định → Name = null, Age = 0 vì không khớp chữ hoa thường.",
           "Bật PropertyNameCaseInsensitive = true → Name = \"An\", Age = 20.",
           "Bước cuối mỗi tab: tùy chọn đã dùng."],
    data="Person An, 20",
    key="Mặc định System.Text.Json phân biệt chữ hoa thường khi đọc; dùng PropertyNameCaseInsensitive hoặc đặt tên khớp.",
    cs="Trong ASP.NET Core, cấu hình web mặc định đã bật không phân biệt hoa thường; ứng dụng console thì không.")
row(6, L, "Thực hành: Serialize đối tượng thành chuỗi JSON", A, K, note="Hoạt động đã có nội dung.")
row(6, L, "Deserialize chuỗi JSON thành đối tượng", A, K, note="Hoạt động đã có nội dung.")
row(6, L, "Bài tập: Xây dựng Hệ thống Quản lý Danh bạ Người dùng", A, K, note="Bài tập nộp, giữ nguyên.")

L = "Các Kỹ Thuật Nâng Cao và Thực Tiễn"
row(6, L, "Xử lý các trường `[NonSerialized]` và `[JsonIgnore]`", R, N,
    goal="Thấy [JsonIgnore] loại một thuộc tính khỏi JSON và hệ quả khi đọc lại.",
    form="2 cột song song: Không có [JsonIgnore] | Có [JsonIgnore]",
    visual="Thẻ User có thuộc tính Password (biểu tượng ổ khóa). JSON hai bên: trái có \"Password\":\"123456\" (tô đỏ cảnh báo lộ dữ liệu), phải không có. Sau Deserialize, Password bên phải = null.",
    comps="H19", tabs="1 tab",
    steps=["User { Name = \"An\", Password = \"123456\" }.",
           "Trái Serialize → JSON chứa mật khẩu.",
           "Phải: [JsonIgnore] trên Password → JSON chỉ có Name.",
           "Deserialize lại: phải Password = null (giá trị mặc định).",
           "Xong: bảng thuộc tính bỏ qua theo từng bộ serializer."],
    data="User An / 123456",
    key="Dùng [JsonIgnore] cho dữ liệu nhạy cảm hoặc dữ liệu tính được; khi đọc lại, thuộc tính đó nhận giá trị mặc định.",
    cs="[NonSerialized] áp dụng cho trường và dùng với BinaryFormatter (đã loại bỏ); System.Text.Json dùng [JsonIgnore]; XmlSerializer dùng [XmlIgnore].")
row(6, L, "Custom Serialization với interface `ISerializable`", R, K,
    note="Không làm mô phỏng: ISerializable gắn với BinaryFormatter (đã loại bỏ). Đề xuất cập nhật học liệu sang JsonConverter tùy chỉnh.")
row(6, L, "Serialization và Versioning của đối tượng", R, N,
    goal="Thấy điều gì xảy ra khi lớp thay đổi nhưng dữ liệu JSON cũ vẫn còn.",
    form="1 cột, 3 tab",
    visual="JSON cũ (v1) bên trái, lớp mới (v2) bên phải; đường nối: thuộc tính khớp (xanh), thuộc tính thừa trong JSON (xám, bị bỏ qua), thuộc tính mới không có dữ liệu (vàng, giá trị mặc định), thuộc tính đổi tên (đỏ, mất dữ liệu).",
    comps="H19", tabs="Thêm thuộc tính | Xóa thuộc tính | Đổi tên thuộc tính",
    steps=["JSON v1: {\"Name\":\"An\",\"Age\":20}.",
           "Thêm: v2 có thêm Email → đọc được, Email = null.",
           "Xóa: v2 bỏ Age → \"Age\" trong JSON bị bỏ qua, không lỗi.",
           "Đổi tên: v2 đổi Name thành FullName → FullName = null (mất dữ liệu).",
           "Sửa: [JsonPropertyName(\"Name\")] trên FullName → đọc đúng \"An\".",
           "Xong."],
    data="JSON v1 cố định",
    key="Thêm thuộc tính thường an toàn; đổi tên phải giữ tên cũ trong JSON bằng [JsonPropertyName].",
    cs="Mặc định System.Text.Json bỏ qua thuộc tính lạ; từ .NET 8 có thể bắt lỗi bằng JsonUnmappedMemberHandling.Disallow.")
row(6, L, "Thực hành: Viết một Custom Serializer đơn giản", A, K, note="Hoạt động đã có nội dung.")
row(6, L, "Xây dựng Lớp Cấu hình Ứng dụng với Tính năng Lưu/Tải tùy chỉnh", A, K, note="Hoạt động đã có nội dung.")

ROWS = [r for r in ROWS if r["lvl"] != K]

# ======================= XUẤT EXCEL =======================
wb = Workbook()
font = "Arial"
H = Font(name=font, bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor="1F5E78")
thin = Side(style="thin", color="C8CCD0")
B = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
PRI_FILL = {N: "D9F2E3", T: "FFF1CC", K: "EEEEEE"}

ws = wb.active
ws.title = "Phan tich"
heads = ["STT", "Chương (Module)", "Bài học (Lesson)", "Học liệu (Activity)", "Loại học liệu", "Mức đề xuất", "Trạng thái",
         "Mục tiêu của mô phỏng", "Dạng giao diện", "Hình minh họa", "Thành phần hình (mã)", "Tab / tình huống",
         "Kịch bản từng bước", "Dữ liệu mẫu", "Điểm chốt (bước cuối)", "Lưu ý kiến thức C#",
         "Đường dẫn file (đề xuất)", "Link GitHub Pages (sau khi đẩy)", "Ghi chú / lý do", "Duyệt (OK / Sửa / Bỏ)", "Ý kiến của bạn"]
ws.append(heads)
for c in ws[1]:
    c.font, c.fill, c.alignment, c.border = H, HF, Alignment(wrap_text=True, vertical="center"), B
ws.row_dimensions[1].height = 32

for i, r in enumerate(ROWS, 1):
    path = url = st = ""
    if r["lvl"] != K:
        path = "%s/%s/%s/%s/index.html" % (COURSE_DIR, slug(MODULES[r["m"] - 1]), slug(r["lesson"]), slug(r["act"]))
        url = BASE_URL + path.replace("index.html", "")
        st = "Done" if os.path.exists(os.path.join(REPO, path)) else "Chưa done"
    ws.append([i, "%d. %s" % (r["m"], MODULES[r["m"] - 1]), r["lesson"], r["act"], r["typ"], r["lvl"], st,
               r["goal"], r["form"], r["visual"], r["comps"], r["tabs"], r["steps"], r["data"], r["key"], r["cs"],
               path, url, r["note"], "", ""])

widths = [5, 24, 26, 32, 13, 13, 11, 34, 20, 42, 12, 26, 70, 26, 34, 40, 36, 36, 30, 13, 24]
for j, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(j)].width = w
last = ws.max_row
for row_ in ws.iter_rows(min_row=2, max_row=last):
    for c in row_:
        c.font = Font(name=font, size=10)
        c.alignment = wrap
        c.border = B
    pri = row_[5].value
    row_[5].fill = PatternFill("solid", fgColor=PRI_FILL[pri])
    row_[5].font = Font(name=font, size=10, bold=(pri == N))
    if row_[6].value:
        done = row_[6].value == "Done"
        row_[6].fill = PatternFill("solid", fgColor="C6EFCE" if done else "FFC7CE")
        row_[6].font = Font(name=font, size=10, bold=True, color="006100" if done else "9C0006")
    for k in (19, 20):
        row_[k].fill = PatternFill("solid", fgColor="FFFFE0")
ws.freeze_panes = "E2"
ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(heads)), last)
dv = DataValidation(type="list", formula1='"OK,Sửa,Bỏ"', allow_blank=True)
ws.add_data_validation(dv); dv.add("T2:T%d" % last)
dv2 = DataValidation(type="list", formula1='"%s,%s,%s"' % (N, T, K), allow_blank=False)
ws.add_data_validation(dv2); dv2.add("F2:F%d" % last)
for pri, col in PRI_FILL.items():
    ws.conditional_formatting.add("F2:F%d" % last, FormulaRule(formula=['$F2="%s"' % pri], fill=PatternFill("solid", fgColor=col)))

# ---- Sheet thư viện hình ----
hs = wb.create_sheet("Thu vien hinh")
hs.append(["Mã", "Thành phần hình minh họa", "Mô tả hình và chuyển động", "Dùng ở STT"])
for c in hs[1]:
    c.font, c.fill, c.border = H, HF, B
for code, name, desc in COMPS:
    used = [str(i) for i, r in enumerate(ROWS, 1) if r["lvl"] != K and code in [x.strip() for x in r["comps"].split(",")]]
    hs.append([code, name, desc, ", ".join(used)])
for row_ in hs.iter_rows(min_row=2):
    for c in row_:
        c.font = Font(name=font, size=10); c.alignment = wrap; c.border = B
    row_[0].font = Font(name=font, size=10, bold=True)
for L_, w in zip("ABCD", [7, 30, 90, 30]):
    hs.column_dimensions[L_].width = w
hs.freeze_panes = "A2"

# ---- Sheet tổng quan ----
s = wb.create_sheet("Tong quan", 0)
s["A1"] = "Phân tích hoạt động mô phỏng – Khóa C# Ứng dụng"
s["A1"].font = Font(name=font, size=14, bold=True)
s["A2"] = "Danh sách %d học liệu cần có hoạt động mô phỏng (đã loại các học liệu không cần: lý thuyết, tóm tắt, trùng ý, bài tập nộp, thảo luận, trắc nghiệm)." % len(ROWS)
s["A2"].font = Font(name=font, size=10, italic=True)
s.merge_cells("A2:G2"); s["A2"].alignment = Alignment(wrap_text=True); s.row_dimensions[2].height = 30
hdr = ["Chương", N, T, "Tổng học liệu", "Đã duyệt OK", "Done", "Chưa done"]
for j, h in enumerate(hdr, 1):
    c = s.cell(row=4, column=j, value=h)
    c.font, c.fill, c.border = H, HF, B
    c.alignment = Alignment(wrap_text=True, vertical="center")
rng = "'Phan tich'!$B$2:$B$%d" % last
prng = "'Phan tich'!$F$2:$F$%d" % last
srng = "'Phan tich'!$G$2:$G$%d" % last
orng = "'Phan tich'!$T$2:$T$%d" % last
for k, name in enumerate(MODULES, 1):
    r_ = 4 + k
    s.cell(row=r_, column=1, value="%d. %s" % (k, name))
    for j, pri in enumerate([N, T], 2):
        s.cell(row=r_, column=j, value='=COUNTIFS(%s,$A%d,%s,"%s")' % (rng, r_, prng, pri))
    s.cell(row=r_, column=4, value="=COUNTIF(%s,$A%d)" % (rng, r_))
    s.cell(row=r_, column=5, value='=COUNTIFS(%s,$A%d,%s,"OK")' % (rng, r_, orng))
    s.cell(row=r_, column=6, value='=COUNTIFS(%s,$A%d,%s,"Done")' % (rng, r_, srng))
    s.cell(row=r_, column=7, value='=COUNTIFS(%s,$A%d,%s,"Chưa done")' % (rng, r_, srng))
tr = 5 + len(MODULES)
s.cell(row=tr, column=1, value="Tổng")
for j in range(2, 8):
    Lc = get_column_letter(j)
    s.cell(row=tr, column=j, value="=SUM(%s5:%s%d)" % (Lc, Lc, tr - 1))
for row_ in s.iter_rows(min_row=5, max_row=tr, max_col=7):
    for c in row_:
        c.font = Font(name=font, size=10, bold=(c.row == tr)); c.border = B
        if c.column > 1:
            c.alignment = Alignment(horizontal="center")
s.column_dimensions["A"].width = 56
for Lc in "BCDEFG":
    s.column_dimensions[Lc].width = 15

notes = [
    ("Cách đọc file", ""),
    (N, "Bài đọc có quá trình chạy hoặc hiện tượng khó hình dung bằng chữ (dịch chuyển phần tử, luồng exception, bộ đệm tệp, phân tích CSV...)."),
    (T, "Có ích nhưng ý chính đã được bài khác hoặc bài đọc truyền tải khá tốt."),
    ("Sheet 'Phan tich'", "Mỗi dòng một học liệu. Các cột từ 'Mục tiêu' tới 'Lưu ý kiến thức C#' là bản mô tả chi tiết để viết mô phỏng. Ô vàng cuối dòng để bạn chọn OK / Sửa / Bỏ và ghi ý kiến."),
    ("Sheet 'Thu vien hinh'", "22 thành phần hình minh họa dùng chung (mã H01–H22). Mỗi mô phỏng kết hợp khung code quen thuộc với 1–3 thành phần hình để đa dạng và trực quan hơn khóa trước."),
    ("", ""),
    ("Nguyên tắc thiết kế", ""),
    ("Đơn giản trước", "Mỗi mô phỏng một ý chính, 4–9 bước mỗi tab, chỉ bấm Bước tiếp / Chạy hết / Làm lại. Không quiz, không thử thách. Tab chỉ để đổi tình huống."),
    ("Hình minh họa", "Ngoài khung Mã nguồn – Console – Bộ nhớ, thêm khung 'Hình minh họa' vẽ bằng HTML/SVG thuần (không thư viện ngoài). Chuyển động nhẹ 200–300 ms; tắt chuyển động khi người dùng bật prefers-reduced-motion."),
    ("Màu sắc thống nhất", "Vàng = đang xét; Xanh lá = thành công/đã xong; Đỏ = lỗi/mất dữ liệu; Xám = bị bỏ qua/không chạy. Có chế độ tối."),
    ("Điện thoại", "Hiển thị tốt ở bề rộng 390px: hai cột xếp chồng, hình tự co, không tràn ngang."),
    ("Sự thật kỹ thuật", "Thông điệp lỗi, mã lỗi biên dịch, số byte, kết quả tính toán phải đúng theo .NET hiện đại; tính bằng code thật trong mô phỏng, không bịa."),
    ("Lưu ý nội dung lỗi thời", "Chương 6 dùng BinaryFormatter/ISerializable: đã bị loại bỏ trong .NET 9. Mô phỏng không minh họa BinaryFormatter; đề xuất cập nhật học liệu."),
    ("Thư mục", "c-sharp-ung-dung/<chuong>/<bai-hoc>/<hoc-lieu>/index.html, đẩy lên GitHub Pages như các khóa trước."),
]
r0 = tr + 2
for k, (a, b) in enumerate(notes):
    ca = s.cell(row=r0 + k, column=1, value=a)
    cb = s.cell(row=r0 + k, column=2, value=b)
    ca.font = Font(name=font, size=10, bold=True)
    cb.font = Font(name=font, size=10)
    s.merge_cells(start_row=r0 + k, start_column=2, end_row=r0 + k, end_column=7)
    cb.alignment = Alignment(wrap_text=True, vertical="top")
    if b and len(b) > 90:
        s.row_dimensions[r0 + k].height = 30
    if a in PRI_FILL:
        ca.fill = PatternFill("solid", fgColor=PRI_FILL[a])

wb.save(OUT)
c = Counter(r["lvl"] for r in ROWS)
print(len(ROWS), dict(c))
print({m: Counter(r["lvl"] for r in ROWS if r["m"] == m) for m in range(1, 7)})
